#!/usr/bin/env python3
"""
verify-deck.py - deterministic gate for generated PPTX decks.

Exit codes
    0  every check PASSed (WARN allowed unless --strict)
    1  at least one FAIL
    2  no FAIL, but something is UNKNOWN or SKIPPED (the deck is NOT verified)

Checks
    V01 file       opens, sha256 recorded
    V02 structure  slide count, slide size
    V03 palette    every colour used on slides is a brand-token colour
    V04 type       fonts are Arial, explicit sizes are on the brand scale
    V05 brand-text tagline spelling, trademark on first mention, navy rail
    V06 logo       official asset (hash match), count, aspect, size, alt text
    V07 alt-text   every picture has real alt text
    V08 leftovers  lorem/TODO/TBD/placeholder text, forbidden patterns
    V09 claims     every number on every slide (and notes) maps to a claim
    V10 sources    each claim's source exists and its quote is really in it
    V11 fit        text-overflow / off-slide estimate (heuristic)
    V12 contrast   WCAG contrast of explicit text colour vs known backdrop
    V13 qa-verdict ls-artifact-qa verdict file bound to THIS build's sha256

Writes <deck>.verify.json (sha256-bound) so a hook / CI can confirm that the
report belongs to the exact file being delivered.

    python verify-deck.py deck.pptx --root . --claims docs/claims.yaml
    python verify-deck.py --self-test
"""

from __future__ import annotations

import argparse
import contextlib
import datetime as dt
import hashlib
import io
import json
import math
import re
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

try:
    import yaml
    from lxml import etree
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
except ImportError as exc:  # pragma: no cover
    print(
        f"missing dependency: {exc.name}\n  pip install python-pptx pyyaml pillow", file=sys.stderr
    )
    sys.exit(2)

VERSION = "1.0.0"
EMU_IN = 914400
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

DEFAULT_PALETTE = {
    "070A40",
    "E63946",
    "00BFFF",
    "F2F2F2",
    "FFFFFF",
    "6B7280",
    "9CA3AF",
    "F7F8F9",
    "121518",
}
DEFAULT_SIZES = {36, 32, 28, 24, 18, 16, 14, 13, 12}
TAGLINE = "ASPIRE. ACT. ACHIEVE."
COMPANY_RE = re.compile(r"lightspeed holdings(?: limited)?", re.I)
TAGLINE_LOOSE_RE = re.compile(r"aspire\W{0,3}act\W{0,3}achieve", re.I)
LEFTOVER_RE = re.compile(
    r"\blorem\b|\bipsum\b|\bTODO\b|\bTBD\b|\bFIXME\b|\bX{3,}\b|\[insert|\bplaceholder\b|<<[^>]*>>",
    re.I,
)

# ---- number handling -------------------------------------------------------
NUMTOK = re.compile(r"\d+(?:,\d{3})*(?:\.\d+)?")
NUM_RE = re.compile(
    r"(?<![\w])[<>\u2265\u2264~\u2248]?\$?\d+(?:,\d{3})*(?:\.\d+)?"
    r"(?:-\$?\d+(?:,\d{3})*(?:\.\d+)?)?"
    r"(?:\s?(?:%|\u00d7|[kmb](?![a-z0-9])|\+))?"
)
IGNORE_RES = [
    re.compile(r"\b\d{4}-\d{2}-\d{2}\b"),  # ISO dates
    re.compile(r"\b(?:19|20)\d{2}\b"),  # years
    re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+"),  # emails
    re.compile(r"(?:https?://|www\.)\S+"),  # urls
]
ORDINAL_RE = re.compile(r"^\s*\d{1,2}[.)]\s+")  # "1. " list prefixes


def norm(s: str) -> str:
    s = (
        s.replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u2212", "-")
        .replace("\u00a0", " ")
    )
    s = re.sub(r"(?<=\d)\s*[x\u00d7](?![A-Za-z0-9])", "\u00d7", s)
    s = re.sub(r"([<>\u2265\u2264~\u2248])\s+(?=[\d$])", r"\1", s)
    return re.sub(r"\s+", " ", s).strip()


def nums(s: str) -> set[str]:
    return {m.replace(",", "") for m in NUMTOK.findall(s)}


# ---- result model ------------------------------------------------------------
@dataclass
class Check:
    id: str
    title: str
    findings: list = field(default_factory=list)
    skipped: str | None = None

    def add(self, level: str, msg: str) -> None:
        self.findings.append((level, msg))

    @property
    def status(self) -> str:
        if self.skipped:
            return "SKIPPED"
        levels = {lv for lv, _ in self.findings}
        for s in ("FAIL", "UNKNOWN", "WARN"):
            if s in levels:
                return s
        return "PASS"


# ---- deck model ----------------------------------------------------------------
@dataclass
class SlideInfo:
    n: int
    slide: object
    top: list
    allshapes: list
    visible: str
    alt: list
    notes: str
    has_chart: bool


def walk(shapes, top=True):
    for sh in shapes:
        yield sh, top
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes, False)


def skippable_placeholder(sh) -> bool:
    try:
        return bool(sh.is_placeholder and sh.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER)
    except (AttributeError, ValueError, TypeError):
        return False


def paragraphs_of(sh) -> list[str]:
    out = []
    if sh.has_text_frame:
        out += [ORDINAL_RE.sub("", p.text) for p in sh.text_frame.paragraphs]
    if getattr(sh, "has_table", False) and sh.has_table:
        for row in sh.table.rows:
            for cell in row.cells:
                out += [ORDINAL_RE.sub("", p.text) for p in cell.text_frame.paragraphs]
    return out


def build_slide_infos(prs) -> list[SlideInfo]:
    infos = []
    for i, slide in enumerate(prs.slides, 1):
        top, allsh, vis, alt, chart = [], [], [], [], False
        for sh, is_top in walk(slide.shapes):
            allsh.append(sh)
            if is_top:
                top.append(sh)
            if skippable_placeholder(sh):
                continue
            vis += paragraphs_of(sh)
            alt += [a for a in sh._element.xpath("./*[1]/p:cNvPr/@descr") if a.strip()]
            chart = chart or bool(getattr(sh, "has_chart", False) and sh.has_chart)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text
        infos.append(SlideInfo(i, slide, top, allsh, "\n".join(vis), alt, notes, chart))
    return infos


# ---- theme / colour helpers -------------------------------------------------------
SCHEME_ALIAS = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}


def load_theme(prs):
    colors, major, minor = {}, None, None
    try:
        root = etree.fromstring(prs.slide_master.part.part_related_by(RT.THEME).blob)
    except (AttributeError, KeyError, ValueError, etree.XMLSyntaxError):
        return colors, major, minor
    ns = {"a": A_NS}
    cs = root.find(".//a:clrScheme", ns)
    if cs is not None:
        for child in cs:
            c = child[0] if len(child) else None
            if c is None:
                continue
            val = c.get("val") if etree.QName(c).localname == "srgbClr" else c.get("lastClr")
            if val:
                colors[etree.QName(child).localname] = val.upper()
    mj = root.find(".//a:fontScheme/a:majorFont/a:latin", ns)
    mn = root.find(".//a:fontScheme/a:minorFont/a:latin", ns)
    return (
        colors,
        (mj.get("typeface") if mj is not None else None),
        (mn.get("typeface") if mn is not None else None),
    )


def resolve_color(c, theme_colors) -> str | None:
    name = etree.QName(c).localname
    if name == "srgbClr":
        return c.get("val", "").upper() or None
    if name == "schemeClr":
        v = c.get("val")
        return theme_colors.get(SCHEME_ALIAS.get(v, v))
    if name == "sysClr":
        return (c.get("lastClr") or "").upper() or None
    return None


def fill_of(el, theme_colors) -> str | None:
    sf = el.xpath("./p:spPr/a:solidFill/*[1]")
    return resolve_color(sf[0], theme_colors) if sf else None


def slide_bg(slide, theme_colors) -> str | None:
    sf = slide._element.xpath("./p:cSld/p:bg/p:bgPr/a:solidFill/*[1]")
    return resolve_color(sf[0], theme_colors) if sf else None


def geom(sh):
    if None in (sh.left, sh.top, sh.width, sh.height):
        return None
    return sh.left, sh.top, sh.width, sh.height


def contains_pt(sh, x, y) -> bool:
    g = geom(sh)
    return bool(g) and g[0] <= x <= g[0] + g[2] and g[1] <= y <= g[1] + g[3]


def backdrop(info: SlideInfo, target, theme_colors) -> str | None:
    """Colour painted behind `target`: own fill, else last filled shape below it containing its centre, else slide bg."""
    own = fill_of(target._element, theme_colors) if target._element.tag.endswith("}sp") else None
    if own:
        return own
    g = geom(target)
    if not g:
        return None
    cx, cy = g[0] + g[2] // 2, g[1] + g[3] // 2
    idx = info.top.index(target)
    for sh in reversed(info.top[:idx]):
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and contains_pt(sh, cx, cy):
            return None  # unknown: image behind
        if sh._element.tag.endswith("}sp") and contains_pt(sh, cx, cy):
            f = fill_of(sh._element, theme_colors)
            if f:
                return f
    return slide_bg(info.slide, theme_colors)


def lum(hexv: str) -> float:
    ch = []
    for i in (0, 2, 4):
        v = int(hexv[i : i + 2], 16) / 255
        ch.append(v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4)
    return 0.2126 * ch[0] + 0.7152 * ch[1] + 0.0722 * ch[2]


def contrast(a: str, b: str) -> float:
    la, lb = sorted((lum(a), lum(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


# ---- claims -------------------------------------------------------------------------
@dataclass
class Licence:
    id: str
    displays: list
    ok: bool
    problems: list
    slides: list | None
    structural: bool
    status: str = "verified"
    warnings: list = field(default_factory=list)


ID_RE = re.compile(r"^[a-z0-9][a-z0-9._-]{1,63}$")
STATUSES = {"verified", "unverified", "placeholder", "retracted"}
KINDS = {"file", "url", "document", "command", "attestation"}


def as_date(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    try:
        return dt.date.fromisoformat(str(v))
    except (ValueError, TypeError):
        return None


def human(s) -> bool:
    return isinstance(s, str) and s.startswith("human:") and len(s) > 6 and "REPLACE" not in s


def claim_problems(c: dict, root: Path, today: dt.date, args) -> tuple[list, list]:
    probs, warns = [], []
    status = c.get("status")
    if status != "verified":
        return [f"status is '{status}', not 'verified'"], warns
    src = c.get("source") or {}
    kind = src.get("kind")
    vdate = as_date(c.get("verified_on"))
    by = c.get("verified_by")
    if kind not in KINDS:
        probs.append(f"source.kind must be one of {sorted(KINDS)}")
    if vdate is None:
        probs.append("verified_on missing or not an ISO date")
    else:
        age = (today - vdate).days
        limit = int(c.get("max_age_days", args.max_age_days))
        if age < 0:
            probs.append("verified_on is in the future")
        elif age > limit:
            probs.append(f"verification expired ({age} days old, limit {limit})")
    if not isinstance(by, str) or not re.match(r"^(human|ci|agent):\S+", by):
        probs.append("verified_by must look like human:<name> | ci:<job> | agent:<id>")
    mechanical = kind == "file" and bool(src.get("quote"))
    if not mechanical and not human(by) and not args.allow_agent_verification:
        probs.append(
            "this source type needs verified_by: human:<name> (agents may not self-verify it)"
        )
    quote = src.get("quote")
    if kind == "file":
        rel = src.get("path")
        p = (root / rel).resolve() if rel else None
        if p is None or root.resolve() not in p.parents or not p.is_file():
            probs.append(f"source file not found inside repo: {rel}")
        elif not quote:
            probs.append("file-sourced claim needs source.quote (exact text that proves it)")
        else:
            text = norm(p.read_text(encoding="utf-8", errors="replace")).casefold()
            if norm(quote).casefold() not in text:
                probs.append(f"quote NOT FOUND in {rel}")
    elif kind in ("url", "document"):
        if not quote:
            probs.append("source.quote required")
        warns.append(f"{kind} source not re-fetched by this tool")
    elif kind == "command":
        for k in ("command", "recorded_output", "recorded_on"):
            if not src.get(k):
                probs.append(f"command source needs source.{k}")
        warns.append("command source not re-run by this tool")
    elif kind == "attestation" and not human(src.get("attested_by")):
        probs.append("attestation needs source.attested_by: human:<name>")
    if quote and not c.get("derivation"):
        qn = nums(str(quote))
        for d in c.get("display", []):
            missing = nums(d) - qn
            if missing:
                probs.append(
                    f"number(s) {sorted(missing)} in display '{d}' not in source.quote (fix quote or add derivation)"
                )
                break
    return probs, warns


def load_claims(
    path: Path, root: Path, today, args, c_claims: Check, c_src: Check
) -> list[Licence]:
    lic: list[Licence] = []
    try:
        doc = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except (OSError, yaml.YAMLError) as e:
        c_claims.add("FAIL", f"cannot parse {path}: {e}")
        return lic
    if doc.get("schema_version") != 1:
        c_claims.add("FAIL", "claims file: schema_version must be 1")
    seen_ids, seen_disp = set(), {}
    for c in doc.get("claims") or []:
        cid = str(c.get("id", ""))
        if not ID_RE.match(cid) or cid in seen_ids:
            c_claims.add("FAIL", f"claim id invalid or duplicate: '{cid}'")
            continue
        seen_ids.add(cid)
        disp = c.get("display")
        if not (
            isinstance(disp, list) and disp and all(isinstance(d, str) and d.strip() for d in disp)
        ):
            c_claims.add("FAIL", f"claim {cid}: display must be a non-empty list of strings")
            continue
        if c.get("status") not in STATUSES:
            c_claims.add("FAIL", f"claim {cid}: status must be one of {sorted(STATUSES)}")
            continue
        for d in disp:
            k = norm(d).casefold()
            if k in seen_disp:
                c_claims.add("FAIL", f"display '{d}' used by both {seen_disp[k]} and {cid}")
            seen_disp[k] = cid
        probs, warns = claim_problems(c, root, today, args)
        lic.append(
            Licence(
                cid,
                [norm(d).casefold() for d in disp],
                not probs,
                probs,
                c.get("slides"),
                False,
                c.get("status"),
                warns,
            )
        )
    for s in doc.get("structural") or []:
        if not s.get("display") or not s.get("reason"):
            c_claims.add("FAIL", "structural entries need display and reason")
            continue
        lic.append(
            Licence(
                "structural:" + s["display"],
                [norm(s["display"]).casefold()],
                True,
                [],
                s.get("slides"),
                True,
            )
        )
    forb = doc.get("forbidden_patterns") or []
    args._forbidden = [
        (re.compile(f["pattern"], re.I), f.get("reason", "forbidden"))
        for f in forb
        if f.get("pattern")
    ]
    return lic


def check_claim_coverage(infos, lic, c: Check, c_src: Check) -> None:
    entries = sorted(((d, e) for e in lic for d in e.displays), key=lambda x: -len(x[0]))
    pats = [
        (re.compile(r"(?<![\w$])(?<!\d[.,])" + re.escape(d) + r"(?!\w)(?!\.?\d)"), e)
        for d, e in entries
    ]
    shown: set[str] = set()
    seen_msgs: set = set()

    def once(chk: Check, level: str, msg: str) -> None:
        key = (chk.id, re.sub(r"\(first shown on slide \d+\)", "", msg))
        if key not in seen_msgs:
            seen_msgs.add(key)
            chk.add(level, msg)

    for info in infos:
        t = norm(info.visible + "\n" + info.notes + "\n" + "\n".join(info.alt)).casefold()
        reported: set[str] = set()
        for pat, e in pats:
            if pat.search(t):
                shown.add(e.id)
                if not e.structural and e.id not in reported:
                    reported.add(e.id)
                    if e.status != "verified":
                        once(
                            c,
                            "FAIL",
                            f"slide {info.n}: claim '{e.id}' is shown but its status is '{e.status}'",
                        )
                    elif not e.ok:
                        for pr in e.problems:
                            once(
                                c_src,
                                "FAIL",
                                f"claim '{e.id}' (first shown on slide {info.n}): {pr}",
                            )
                    elif e.slides and info.n not in e.slides:
                        once(
                            c,
                            "FAIL",
                            f"slide {info.n}: claim '{e.id}' is restricted to slides {e.slides}",
                        )
                    else:
                        for w in e.warnings:
                            once(c_src, "WARN", f"claim {e.id}: {w}")
            t = pat.sub(lambda m: "\x00" * len(m.group()), t)
        for rx in IGNORE_RES:
            t = rx.sub(lambda m: "\x01" * len(m.group()), t)
        for m in NUM_RE.finditer(t):
            ctx = re.sub(r"[\x00\x01]", "_", t[max(0, m.start() - 28) : m.end() + 28])
            once(c, "FAIL", f"slide {info.n}: unsourced number '{m.group().strip()}' ...{ctx}...")
        if info.has_chart:
            c.add("UNKNOWN", f"slide {info.n}: contains a chart; chart data is not claim-checked")
    for e in lic:
        if not e.structural and e.id not in shown:
            note = "" if e.ok else f" and is not usable ({'; '.join(e.problems)})"
            c_src.add("WARN", f"claim '{e.id}' is not shown on any slide{note}")


# ---- individual checks ----------------------------------------------------------------------
def tok_value(tokens: dict, dotted: str, default=None):
    cur = tokens
    for k in dotted.split("."):
        if not isinstance(cur, dict) or k not in cur:
            return default
        cur = cur[k]
    return cur


def check_structure(prs, infos, args, tokens, c: Check) -> None:
    if len(infos) != args.slides:
        c.add("FAIL", f"slide count {len(infos)} != required {args.slides}")
    w, h = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
    sw = tok_value(tokens, "layout.slide.width", "10in")
    sh_ = tok_value(tokens, "layout.slide.height", "7.5in")
    ew, eh = float(str(sw).replace("in", "")), float(str(sh_).replace("in", ""))
    if abs(w - ew) > 0.01 or abs(h - eh) > 0.01:
        c.add("FAIL", f"slide size {w:.2f}x{h:.2f}in != brand {ew}x{eh}in")


def check_palette(infos, theme, palette, c: Check) -> None:
    explicit = "./p:spPr/a:solidFill|./p:spPr/a:noFill|./p:spPr/a:gradFill|./p:spPr/a:blipFill|./p:spPr/a:pattFill"
    explicit_ln = "./p:spPr/a:ln/a:solidFill|./p:spPr/a:ln/a:noFill"
    for info in infos:
        bad = set()
        el = info.slide._element
        for v in el.xpath(".//a:srgbClr/@val"):
            if v.upper() not in palette:
                bad.add(v.upper())
        for v in el.xpath(".//a:schemeClr[not(ancestor::p:style)]/@val"):
            r = theme.get(SCHEME_ALIAS.get(v, v))
            if r is None:
                c.add("UNKNOWN", f"slide {info.n}: scheme colour '{v}' not resolvable")
            elif r not in palette:
                bad.add(f"{r}(scheme:{v})")
        for v in el.xpath(".//a:prstClr/@val"):
            bad.add(f"preset:{v}")
        # shapes that rely on theme style refs (no explicit fill / line) inherit theme colours
        for sp in el.xpath(".//p:sp[p:style]"):
            for ref, has_explicit in (("a:fillRef", explicit), ("a:lnRef", explicit_ln)):
                r = sp.xpath(f"./p:style/{ref}")
                if r and r[0].get("idx") not in (None, "0") and not sp.xpath(has_explicit):
                    kids = r[0].xpath("./*[1]")
                    hexv = resolve_color(kids[0], theme) if kids else None
                    if hexv and hexv not in palette:
                        bad.add(f"{hexv}(theme style {ref[2:]})")
        if bad:
            c.add("FAIL", f"slide {info.n}: off-palette colours {sorted(bad)}")


def check_type(infos, theme_fonts, sizes, c: Check) -> None:
    major, minor = theme_fonts
    inherited_font, inherited_size = [], []
    for info in infos:
        for r in info.slide._element.xpath(".//a:r"):
            txt = "".join(r.xpath("./a:t/text()")).strip()
            if not txt:
                continue
            face = r.xpath("./a:rPr/a:latin/@typeface")
            if face:
                f = {"+mn-lt": minor, "+mj-lt": major}.get(face[0], face[0])
                if (f or "").lower() != "arial":
                    c.add("FAIL", f"slide {info.n}: font '{f}' on text '{txt[:30]}' (brand: Arial)")
            else:
                inherited_font.append(info.n)
            sz = r.xpath("./a:rPr/@sz")
            if sz:
                pt = int(sz[0]) / 100
                if pt not in sizes:
                    c.add(
                        "FAIL", f"slide {info.n}: size {pt:g}pt not on brand scale {sorted(sizes)}"
                    )
            else:
                inherited_size.append(info.n)
    if inherited_font and any((f or "").lower() != "arial" for f in (major, minor)):
        c.add(
            "FAIL",
            f"{len(inherited_font)} runs inherit theme font ({major}/{minor}) which is not Arial "
            f"(slides {sorted(set(inherited_font))})",
        )
    if inherited_size:
        c.add(
            "UNKNOWN",
            f"{len(inherited_size)} runs inherit font size (not set explicitly) on slides "
            f"{sorted(set(inherited_size))}; set sz on every run",
        )


def check_brand_text(infos, theme, palette, args, c: Check) -> None:
    found_tagline = False
    for info in infos:
        for m in TAGLINE_LOOSE_RE.finditer(info.visible):
            span = info.visible[m.start() : m.start() + len(TAGLINE)]
            if span == TAGLINE:
                found_tagline = True
            else:
                c.add(
                    "FAIL", f"slide {info.n}: tagline spelled '{m.group()}' (must be '{TAGLINE}')"
                )
    if not found_tagline and args.require_tagline:
        c.add("FAIL", f"exact tagline '{TAGLINE}' not found on any slide")
    first = None
    for info in infos:
        m = COMPANY_RE.search(info.visible)
        if m:
            first = (info.n, info.visible[m.end() : m.end() + 2])
            break
    if first is None:
        c.add("FAIL", "company name 'LightSpeed Holdings Limited' not found on any slide")
    elif not first[1].lstrip().startswith("\u2122"):
        c.add(
            "FAIL", f"first mention of company name (slide {first[0]}) lacks the trademark symbol"
        )
    navy = (tok_value(args._tokens, "color.navy.value", "#070A40")).lstrip("#").upper()
    exempt = {int(x) for x in args.rail_exempt.split(",") if x.strip()}
    for info in infos:
        if info.n in exempt:
            continue
        ok = False
        for sh in info.top:
            g = geom(sh)
            if (
                g
                and sh._element.tag.endswith("}sp")
                and g[0] <= 0.02 * EMU_IN
                and abs(g[2] / EMU_IN - 0.15) <= 0.03
                and g[3]
                >= info.slide.part.package.presentation_part.presentation.slide_height * 0.95
                and fill_of(sh._element, theme) == navy
            ):
                ok = True
                break
        if not ok:
            c.add(
                "FAIL",
                f"slide {info.n}: navy 0.15in left rail missing (use --rail-exempt for cover/closing)",
            )


def check_logo_and_alt(infos, theme, root, args, c_logo: Check, c_alt: Check) -> None:
    official: dict[str, str] = {}
    for d in args.logo_dir or ["brand/logo", "static/brand/logo"]:
        base = root / d
        if base.is_dir():
            for p in base.rglob("*"):
                if p.suffix.lower() in (".png", ".jpg", ".jpeg"):
                    official[hashlib.sha256(p.read_bytes()).hexdigest()] = p.relative_to(
                        root
                    ).as_posix()
    from PIL import Image  # python-pptx already depends on Pillow

    found: list[tuple[int, str]] = []
    for info in infos:
        for sh in info.allshapes:
            if sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            descr = (sh._element.xpath("./p:nvPicPr/p:cNvPr/@descr") or [""])[0].strip()
            name = (sh._element.xpath("./p:nvPicPr/p:cNvPr/@name") or [""])[0]
            if not descr:
                c_alt.add("FAIL", f"slide {info.n}: picture '{name}' has no alt text")
            elif re.search(r"\.(png|jpe?g|gif)$", descr, re.I):
                c_alt.add("FAIL", f"slide {info.n}: alt text is a filename ('{descr}')")
            sha = hashlib.sha256(sh.image.blob).hexdigest()
            src = official.get(sha)
            if not src:
                if "logo" in (name + descr).lower():
                    c_logo.add(
                        "FAIL",
                        f"slide {info.n}: picture '{name}' looks like a logo but is not byte-identical to any official asset",
                    )
                continue
            found.append((info.n, src))
            is_top = sh in info.top
            if sh.rotation:
                c_logo.add("FAIL", f"slide {info.n}: logo rotated {sh.rotation} deg")
            if any((sh.crop_left, sh.crop_right, sh.crop_top, sh.crop_bottom)):
                c_logo.add("FAIL", f"slide {info.n}: logo is cropped")
            iw, ih = Image.open(io.BytesIO(sh.image.blob)).size
            if sh.width and sh.height and abs((sh.width / sh.height) / (iw / ih) - 1) > 0.02:
                c_logo.add(
                    "FAIL",
                    f"slide {info.n}: logo aspect ratio distorted "
                    f"({sh.width / sh.height:.3f} vs source {iw / ih:.3f})",
                )
            min_in = args.min_icon_in if "logo-mark" in src else args.min_logo_in
            if sh.width and sh.width / EMU_IN < min_in:
                c_logo.add(
                    "FAIL",
                    f"slide {info.n}: logo {sh.width / EMU_IN:.2f}in wide < minimum {min_in:.2f}in",
                )
            if is_top:
                bg = backdrop(info, sh, theme)
                if bg is None:
                    c_logo.add("UNKNOWN", f"slide {info.n}: logo backdrop colour unknown")
                elif bg != "FFFFFF" and "transparent" not in Path(src).name.lower():
                    c_logo.add(
                        "FAIL",
                        f"slide {info.n}: non-transparent logo '{src}' on non-white backdrop #{bg}",
                    )
                g = geom(sh)
                pad = int(args.clearspace_ratio * g[3])
                zone = (g[0] - pad, g[1] - pad, g[0] + g[2] + pad, g[1] + g[3] + pad)
                sw_, sh_h = (
                    info.slide.part.package.presentation_part.presentation.slide_width,
                    info.slide.part.package.presentation_part.presentation.slide_height,
                )
                for o in info.top:
                    og = geom(o)
                    if o is sh or not og:
                        continue
                    if og[2] * og[3] > 0.8 * sw_ * sh_h:
                        continue
                    if (
                        og[0] <= g[0]
                        and og[1] <= g[1]
                        and og[0] + og[2] >= g[0] + g[2]
                        and og[1] + og[3] >= g[1] + g[3]
                    ):
                        continue  # backdrop panel containing the logo
                    if (
                        og[0] < zone[2]
                        and og[0] + og[2] > zone[0]
                        and og[1] < zone[3]
                        and og[1] + og[3] > zone[1]
                    ):
                        c_logo.add(
                            "WARN",
                            f"slide {info.n}: '{o.name}' intrudes into approximate logo clear space "
                            f"(ratio {args.clearspace_ratio}; real rule is 1x 'L' height - measure by eye)",
                        )
            else:
                c_logo.add(
                    "UNKNOWN", f"slide {info.n}: logo inside a group; geometry checks skipped"
                )
    if not official:
        c_logo.add(
            "UNKNOWN", "no official logo assets found (--logo-dir); cannot prove logos are official"
        )
    if args.logo_slides:
        want = sorted(int(x) for x in args.logo_slides.split(","))
        got = sorted(n for n, _ in found)
        if got != want:
            c_logo.add("FAIL", f"logo on slides {got}, required exactly {want}")
    elif len(found) != args.logo_count:
        c_logo.add(
            "FAIL",
            f"{len(found)} official logo placements, required exactly {args.logo_count} "
            f"(on slides {[n for n, _ in found]})",
        )


def check_leftovers(infos, args, c: Check) -> None:
    extra = getattr(args, "_forbidden", [])
    for info in infos:
        blob = info.visible + "\n" + info.notes
        for w in sorted({m.group().lower() for m in LEFTOVER_RE.finditer(blob)}):
            c.add("FAIL", f"slide {info.n}: leftover/placeholder text '{w}'")
        for rx, why in extra:
            for w in sorted({m.group() for m in rx.finditer(blob)}):
                c.add("FAIL", f"slide {info.n}: forbidden text '{w}' ({why})")


def check_fit(prs, infos, theme, args, c: Check) -> None:
    c.add(
        "WARN",
        "heuristic estimate (Arial ~0.5em/char, 1.2 line height); the rendered-image review in V13 is authoritative",
    ) if args.show_heuristic_note else None
    sw, shh = prs.slide_width, prs.slide_height
    for info in infos:
        for sh in info.top:
            g = geom(sh)
            if not g:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    c.add(
                        "UNKNOWN",
                        f"slide {info.n}: '{sh.name}' inherits geometry from layout; cannot check fit",
                    )
                continue
            tol = int(0.01 * EMU_IN)
            if g[0] < -tol or g[1] < -tol or g[0] + g[2] > sw + tol or g[1] + g[3] > shh + tol:
                c.add("FAIL", f"slide {info.n}: '{sh.name}' extends beyond the slide")
            if not (sh.has_text_frame and sh.text_frame.text.strip()):
                continue
            body = sh._element.xpath("./p:txBody/a:bodyPr")
            if body and body[0].get("wrap") == "none":
                continue
            if body and body[0].xpath("./a:normAutofit"):
                c.add(
                    "UNKNOWN",
                    f"slide {info.n}: '{sh.name}' uses shrink-on-overflow; fit not simulated",
                )
                continue
            tf = sh.text_frame
            wid = (g[2] - (tf.margin_left or 0) - (tf.margin_right or 0)) / EMU_IN
            hei = (g[3] - (tf.margin_top or 0) - (tf.margin_bottom or 0)) / EMU_IN
            need = 0.0
            for p in tf.paragraphs:
                szs = [int(s) / 100 for s in p._p.xpath(".//a:rPr/@sz")]
                pt = max(szs) if szs else 18.0
                chars = len(p.text)
                lines = max(1, math.ceil(chars * 0.5 * pt / 72 / max(wid, 0.1)))
                need += lines * pt * 1.2 / 72
            ratio = need / max(hei, 0.05)
            if body and body[0].xpath("./a:spAutoFit"):
                if g[1] + need * EMU_IN > shh:
                    c.add(
                        "FAIL",
                        f"slide {info.n}: auto-growing '{sh.name}' would run past slide bottom",
                    )
            elif ratio > 1.15:
                c.add(
                    "FAIL",
                    f"slide {info.n}: '{sh.name}' text needs ~{need:.2f}in, box has {hei:.2f}in (x{ratio:.2f})",
                )
            elif ratio > 1.0:
                c.add("WARN", f"slide {info.n}: '{sh.name}' text nearly overflows (x{ratio:.2f})")


def check_contrast(infos, theme, c: Check) -> None:
    unknown_runs: dict[int, int] = {}
    for info in infos:
        for sh in info.top:
            if not (sh.has_text_frame and sh.text_frame.text.strip()):
                continue
            bg = backdrop(info, sh, theme)
            for r in sh._element.xpath(".//a:r"):
                if not "".join(r.xpath("./a:t/text()")).strip():
                    continue
                fc = r.xpath("./a:rPr/a:solidFill/*[1]")
                fg = resolve_color(fc[0], theme) if fc else None
                if not fg or not bg:
                    unknown_runs[info.n] = unknown_runs.get(info.n, 0) + 1
                    continue
                sz = (r.xpath("./a:rPr/@sz") or ["0"])[0]
                pt = int(sz) / 100
                bold = (r.xpath("./a:rPr/@b") or ["0"])[0] in ("1", "true")
                need = 3.0 if (pt >= 18 or (pt >= 14 and bold)) else 4.5
                got = contrast(fg, bg)
                if got < need:
                    txt = "".join(r.xpath("./a:t/text()"))[:28]
                    c.add(
                        "FAIL",
                        f"slide {info.n}: #{fg} on #{bg} = {got:.2f}:1 < {need}:1 for '{txt}' ({pt:g}pt)",
                    )
    if unknown_runs:
        c.add(
            "UNKNOWN",
            "text colour or backdrop not explicit (contrast not computable): "
            + ", ".join(f"slide {k}: {v} runs" for k, v in sorted(unknown_runs.items())),
        )


def check_qa(args, sha: str, n_slides: int, c: Check) -> None:
    if args.skip_qa:
        c.skipped = "--skip-qa given (deck is NOT gated)"
        return
    p = Path(args.qa_verdict)
    if not p.is_file():
        c.add(
            "FAIL",
            f"no QA verdict at {p} (ls-artifact-qa must write it after rendering the slides)",
        )
        return
    try:
        v = json.loads(p.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e:
        c.add("FAIL", f"verdict file unreadable: {e}")
        return
    if v.get("verdict") != "APPROVE":
        c.add("FAIL", f"verdict is '{v.get('verdict')}', not APPROVE")
    if (v.get("artifact") or {}).get("sha256") != sha:
        c.add(
            "FAIL",
            "verdict was issued for a different build (artifact.sha256 does not match this deck)",
        )
    got = {x.get("pass"): x.get("result") for x in v.get("checks", [])}
    for need in ("visual", "brand", "ux", "accessibility", "content"):
        if got.get(need) != "PASS":
            c.add("FAIL", f"QA pass '{need}' is '{got.get(need)}', must be PASS")
    if not v.get("generated_by"):
        c.add("FAIL", "verdict.generated_by missing")
    ev = v.get("evidence") or []
    if len(ev) < n_slides:
        c.add(
            "FAIL", f"{len(ev)} evidence renders listed, need at least {n_slides} (one per slide)"
        )
    from PIL import Image

    for e in ev:
        ep = (p.parent / e.get("path", "")).resolve()
        if not ep.is_file():
            c.add("FAIL", f"evidence missing: {e.get('path')}")
            continue
        if hashlib.sha256(ep.read_bytes()).hexdigest() != e.get("sha256"):
            c.add("FAIL", f"evidence hash mismatch: {e.get('path')}")
            continue
        try:
            if Image.open(ep).size[0] < 640:
                c.add("FAIL", f"evidence image too small to inspect: {e.get('path')}")
        except (OSError, Image.UnidentifiedImageError, ValueError):
            c.add("FAIL", f"evidence is not an image: {e.get('path')}")


# ---- driver ----------------------------------------------------------------------------------------
def run(args) -> int:
    root = Path(args.root).resolve()
    deck = Path(args.deck).resolve()
    checks = {
        k: Check(k, t)
        for k, t in [
            ("V01", "file"),
            ("V02", "structure"),
            ("V03", "palette"),
            ("V04", "type"),
            ("V05", "brand-text"),
            ("V06", "logo"),
            ("V07", "alt-text"),
            ("V08", "leftovers"),
            ("V09", "claims"),
            ("V10", "sources"),
            ("V11", "fit"),
            ("V12", "contrast"),
            ("V13", "qa-verdict"),
        ]
    }
    today = as_date(args.today) if args.today else dt.date.today()
    sha = ""
    try:
        data = deck.read_bytes()
        sha = hashlib.sha256(data).hexdigest()
        prs = Presentation(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001 - malformed deck must FAIL the gate, not crash the run
        checks["V01"].add("FAIL", f"cannot open {deck}: {e}")
        return finish(args, deck, sha, 0, checks, None)

    tokens: dict = {}
    tp = root / args.tokens
    if tp.is_file():
        tokens = json.loads(tp.read_text(encoding="utf-8"))
        palette = {
            str(v["value"]).lstrip("#").upper()
            for v in tok_value(tokens, "color", {}).values()
            if "value" in v
        }
        sizes = {
            float(str(v).replace("pt", ""))
            for v in tok_value(tokens, "typography.sizes", {}).values()
        } or DEFAULT_SIZES
    else:
        palette, sizes = DEFAULT_PALETTE, DEFAULT_SIZES
        checks["V03"].add("WARN", f"tokens file not found at {tp}; using built-in palette")
    args._tokens = tokens
    for lim in ("fullLogoScreen", "iconOnlyScreen"):
        m = re.search(r"(\d+)\s*px", str(tok_value(tokens, f"brand.logo.minimumSizes.{lim}", "")))
        if m:
            if lim == "fullLogoScreen" and args.min_logo_in is None:
                args.min_logo_in = int(m.group(1)) / 96
            if lim == "iconOnlyScreen" and args.min_icon_in is None:
                args.min_icon_in = int(m.group(1)) / 96
    args.min_logo_in = args.min_logo_in or 120 / 96
    args.min_icon_in = args.min_icon_in or 32 / 96

    theme, major, minor = load_theme(prs)
    infos = build_slide_infos(prs)
    checks["V01"].findings.append(("INFO", f"sha256 {sha}"))

    check_structure(prs, infos, args, tokens, checks["V02"])
    check_palette(infos, theme, palette, checks["V03"])
    check_type(infos, (major, minor), sizes, checks["V04"])
    check_brand_text(infos, theme, palette, args, checks["V05"])
    check_logo_and_alt(infos, theme, root, args, checks["V06"], checks["V07"])
    cp = root / args.claims
    if cp.is_file():
        lic = load_claims(cp, root, today, args, checks["V09"], checks["V10"])
        check_claim_coverage(infos, lic, checks["V09"], checks["V10"])
    else:
        checks["V09"].add(
            "FAIL", f"claims file not found: {cp} (every number must be backed by a claim)"
        )
        checks["V10"].skipped = "no claims file"
    check_leftovers(infos, args, checks["V08"])
    check_fit(prs, infos, theme, args, checks["V11"])
    check_contrast(infos, theme, checks["V12"])
    check_qa(args, sha, len(infos), checks["V13"])
    return finish(args, deck, sha, len(infos), checks, cp)


def ranges(ns: list[int]) -> str:
    ns, out, i = sorted(set(ns)), [], 0
    while i < len(ns):
        j = i
        while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
            j += 1
        out.append(str(ns[i]) if i == j else f"{ns[i]}-{ns[j]}")
        i = j + 1
    return ",".join(out)


def group_findings(findings):
    """Collapse identical per-slide messages: 'slide 2: x' + 'slide 3: x' -> 'slides 2-3: x'."""
    order, groups = [], {}
    for lv, msg in findings:
        if lv == "INFO":
            continue
        m = re.match(r"^slide (\d+): (.*)$", msg, re.S)
        key = (lv, m.group(2) if m else msg, bool(m))
        if key not in groups:
            groups[key] = []
            order.append(key)
        if m:
            groups[key].append(int(m.group(1)))
    for key in order:
        lv, body, per_slide = key
        if per_slide:
            ns = groups[key]
            yield lv, f"slide{'s' if len(set(ns)) > 1 else ''} {ranges(ns)}: {body}"
        else:
            yield lv, body


def finish(args, deck: Path, sha: str, n: int, checks: dict, claims_path) -> int:
    relax = {x.strip() for x in (args.allow_unknown or "").split(",") if x.strip()}
    for k, c in checks.items():
        if k in relax:
            c.findings = [("WARN" if lv == "UNKNOWN" else lv, m) for lv, m in c.findings]
    statuses = {k: c.status for k, c in checks.items()}
    vals = set(statuses.values())
    code = (
        1
        if "FAIL" in vals or (args.strict and "WARN" in vals)
        else 2
        if ({"UNKNOWN", "SKIPPED"} & vals)
        else 0
    )
    for c in checks.values():
        print(f"{c.id}  {c.status:<8} {c.title}")
    for c in checks.values():
        for lv, msg in group_findings(c.findings):
            print(f"  [{c.id}] {lv}: {msg}")
    verdict = {0: "VERIFIED", 1: "REJECTED", 2: "INCOMPLETE (not verified)"}[code]
    print(f"\n{verdict}  exit={code}  slides={n}  sha256={sha[:16]}...")
    report = {
        "schema_version": 1,
        "tool": f"verify-deck.py {VERSION}",
        "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "artifact": {"path": str(deck), "sha256": sha, "slides": n},
        "claims_sha256": hashlib.sha256(claims_path.read_bytes()).hexdigest()
        if claims_path and claims_path.is_file()
        else None,
        "exit_code": code,
        "verdict": verdict,
        "checks": [
            {
                "id": c.id,
                "title": c.title,
                "status": c.status,
                "findings": [{"level": lv, "message": m} for lv, m in c.findings if lv != "INFO"],
                "skipped": c.skipped,
            }
            for c in checks.values()
        ],
    }
    out = Path(args.report) if args.report else deck.with_suffix(".verify.json")
    if sha:
        out.write_text(json.dumps(report, indent=2), encoding="utf-8")
        print(f"report: {out}")
    return code


def build_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("deck", nargs="?")
    ap.add_argument("--root", default=".", help="repo root; all relative paths resolve here")
    ap.add_argument("--claims", default="docs/claims.yaml")
    ap.add_argument("--tokens", default="brand/tokens/brand-tokens.json")
    ap.add_argument(
        "--logo-dir", action="append", help="official logo dir (repeatable; default brand/logo)"
    )
    ap.add_argument("--slides", type=int, default=13)
    ap.add_argument(
        "--logo-count", type=int, default=1, help="exact number of official logo placements"
    )
    ap.add_argument("--logo-slides", help="comma list; overrides --logo-count (e.g. 1,13)")
    ap.add_argument("--min-logo-in", type=float)
    ap.add_argument("--min-icon-in", type=float)
    ap.add_argument("--clearspace-ratio", type=float, default=0.15)
    ap.add_argument("--rail-exempt", default="1,13", help="slides without the navy rail")
    ap.add_argument("--no-require-tagline", dest="require_tagline", action="store_false")
    ap.add_argument("--qa-verdict")
    ap.add_argument(
        "--skip-qa", action="store_true", help="dev only: result can never exceed exit 2"
    )
    ap.add_argument("--max-age-days", type=int, default=90)
    ap.add_argument("--allow-agent-verification", action="store_true")
    ap.add_argument(
        "--allow-unknown", help="comma list of check ids whose UNKNOWN becomes WARN, e.g. V11,V12"
    )
    ap.add_argument("--strict", action="store_true", help="WARN also fails")
    ap.add_argument("--report", help="path for JSON report (default <deck>.verify.json)")
    ap.add_argument("--today", help="override today's date (tests)")
    ap.add_argument("--show-heuristic-note", action="store_true", help=argparse.SUPPRESS)
    ap.add_argument(
        "--self-test",
        action="store_true",
        help="build good/bad fixture decks and assert the gate behaves",
    )
    return ap


def main(argv=None) -> int:
    with contextlib.suppress(Exception):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    args = build_parser().parse_args(argv)
    if args.self_test:
        return self_test()
    if not args.deck:
        print("usage: verify-deck.py DECK.pptx [options]", file=sys.stderr)
        return 2
    if not args.qa_verdict:
        args.qa_verdict = str(Path(args.deck).with_suffix(".qa-verdict.json"))
    return run(args)


# ---- self test: the gate must be RED when it should be ----------------------------------------------------
def self_test() -> int:
    from PIL import Image
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Inches, Pt

    tmp = Path(tempfile.mkdtemp(prefix="verify-deck-selftest-"))
    (tmp / "brand/tokens").mkdir(parents=True)
    (tmp / "docs").mkdir()
    (tmp / "brand/logo").mkdir(parents=True)
    tokens = {
        "brand": {
            "logo": {
                "minimumSizes": {"fullLogoScreen": "120px wide", "iconOnlyScreen": "32px wide"}
            }
        },
        "color": {
            k: {"value": v}
            for k, v in zip(
                ["navy", "red", "cyan", "white", "grey-light"],
                ["#070A40", "#E63946", "#00BFFF", "#FFFFFF", "#F2F2F2"],
                strict=True,
            )
        },
        "typography": {"sizes": {"a": "36pt", "b": "28pt", "c": "16pt", "d": "14pt"}},
        "layout": {"slide": {"width": "10in", "height": "7.5in"}},
    }
    (tmp / "brand/tokens/brand-tokens.json").write_text(json.dumps(tokens))
    logo = tmp / "brand/logo/logo-transparent-selftest.png"
    Image.new("RGBA", (400, 100), (7, 10, 64, 255)).save(logo)
    rogue = tmp / "rogue.png"
    Image.new("RGBA", (400, 100), (200, 0, 0, 255)).save(rogue)
    (tmp / "company-registry.yaml").write_text("agents: 90\ndepartments: 20\n")
    today = dt.date.today().isoformat()
    (tmp / "docs/claims.yaml").write_text(f"""
schema_version: 1
claims:
  - id: agents-total
    display: ["90 agents"]
    status: verified
    verified_on: {today}
    verified_by: ci:selftest
    source: {{kind: file, path: company-registry.yaml, quote: "agents: 90"}}
  - id: atc-rate
    display: ["ATC >90%"]
    status: verified
    verified_on: {today}
    verified_by: ci:selftest
    source: {{kind: file, path: company-registry.yaml, quote: "atc: 0.9"}}
  - id: tam
    display: ["$65B"]
    status: placeholder
structural: []
forbidden_patterns:
  - {{pattern: "Fortune \\\\d+", reason: "no unverifiable client-tier claims"}}
""")

    def rect(slide, x, y, w, h, hexv):
        s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(hexv)
        s.line.fill.background()
        return s

    def text(slide, x, y, w, h, s, pt, hexv, font="Arial", bold=False):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        r = tf.paragraphs[0].add_run()
        r.text = s
        r.font.size, r.font.name, r.font.bold = Pt(pt), font, bold
        r.font.color.rgb = RGBColor.from_string(hexv)

    def make(path: Path, bad: bool) -> None:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
        n = 12 if bad else 13
        for i in range(1, n + 1):
            s = prs.slides.add_slide(prs.slide_layouts[6])
            if i == 1:
                rect(s, 0, 0, 10, 7.5, "070A40")
                text(
                    s, 0.8, 2.5, 8, 1, "LightSpeed Holdings Limited\u2122", 36, "FFFFFF", bold=True
                )
                text(s, 0.8, 3.8, 8, 0.6, "Aspire. Act. Achieve" if bad else TAGLINE, 16, "00BFFF")
                pic = s.shapes.add_picture(
                    str(rogue if bad else logo), Inches(0.8), Inches(0.6), width=Inches(2)
                )
                pic._element.xpath("./p:nvPicPr/p:cNvPr")[0].set("name", "Company logo")
                if not bad:
                    pic._element.xpath("./p:nvPicPr/p:cNvPr")[0].set(
                        "descr", "LightSpeed Holdings logo"
                    )
            else:
                rect(s, 0, 0, 10, 7.5, "FFFFFF")
                if not (bad and i == 3):
                    rect(s, 0, 0, 0.15, 7.5, "070A40")
                text(s, 0.6, 0.4, 8.8, 0.8, "Section title", 28, "070A40", bold=True)
                body = (
                    "The registry defines 90 agents."
                    if not bad
                    else "ATC >90% and a $65B market. Seed ask $500K. Trusted by Fortune 500 firms. "
                    + "Lorem " * 60
                )
                text(
                    s,
                    0.6,
                    1.5,
                    8.8,
                    1.0,
                    body,
                    14,
                    "070A40" if not bad else "00FF00",
                    font="Arial" if not bad else "Calibri",
                )
        prs.save(str(path))

    good, bad = tmp / "good.pptx", tmp / "bad.pptx"
    make(good, False)
    make(bad, True)

    def verdict_for(deck: Path) -> None:
        ev = []
        (tmp / "renders").mkdir(exist_ok=True)
        for i in range(1, 14):
            p = tmp / "renders" / f"slide-{i}.png"
            Image.new("RGB", (800, 600), (255, 255, 255)).save(p)
            ev.append(
                {
                    "path": f"renders/slide-{i}.png",
                    "sha256": hashlib.sha256(p.read_bytes()).hexdigest(),
                }
            )
        v = {
            "schema_version": 1,
            "verdict": "APPROVE",
            "generated_by": "selftest",
            "artifact": {
                "path": deck.name,
                "sha256": hashlib.sha256(deck.read_bytes()).hexdigest(),
            },
            "checks": [
                {"pass": k, "result": "PASS"}
                for k in ("visual", "brand", "ux", "accessibility", "content")
            ],
            "evidence": ev,
        }
        deck.with_suffix(".qa-verdict.json").write_text(json.dumps(v))

    verdict_for(good)
    verdict_for(bad)
    base = ["--root", str(tmp)]
    print("=== GOOD deck (expect exit 0) ===")
    rc_good = main([str(good)] + base)
    print("\n=== BAD deck (expect exit 1 with specific failures) ===")
    rc_bad = main([str(bad)] + base)
    rep = json.loads(bad.with_suffix(".verify.json").read_text())
    failed = {c["id"] for c in rep["checks"] if c["status"] == "FAIL"}
    blob = json.dumps(rep)
    print("\n=== STALE VERDICT (good deck edited after QA; expect exit 1) ===")
    prs = Presentation(str(good))
    prs.core_properties.title = "tampered"
    prs.save(str(good))
    rc_stale = main([str(good)] + base)
    print("\n=== QA SKIPPED (expect exit 2: incomplete, never VERIFIED) ===")
    rc_skip = main([str(good), "--skip-qa"] + base)

    expect_ids = {"V02", "V03", "V04", "V05", "V06", "V07", "V08", "V09", "V10", "V12"}
    expect_text = [
        "off-palette",
        "tagline spelled",
        "unsourced number '$500k'",
        "quote NOT FOUND",
        "claim 'tam' is shown but its status is 'placeholder'",
        "not byte-identical",
        "alt text is a filename",
        "navy 0.15in left rail missing",
        "forbidden text",
        "slide count 12",
    ]
    ok = (
        rc_good == 0
        and rc_bad == 1
        and rc_stale == 1
        and rc_skip == 2
        and expect_ids <= failed
        and all(t in blob for t in expect_text)
    )
    print("\nSELF-TEST:", "PASS" if ok else "FAIL")
    if not ok:
        print(
            "  rc:",
            rc_good,
            rc_bad,
            rc_stale,
            rc_skip,
            "| missing ids:",
            sorted(expect_ids - failed),
            "| missing text:",
            [t for t in expect_text if t not in blob],
        )
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
