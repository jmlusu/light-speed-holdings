"""Verify the AI Venture Studio client deck against the 10 QA acceptance checks.

Usage:
    uv run --with python-pptx python scripts/presentation/verify-venture-studio-deck.py [deck.pptx]

Checks:
 1. Exactly 13 slides.
 2. Slide size 10 x 7.5 in.
 3. Zero non-Arial runs.
 4. Zero off-palette RGB values.
 5. Navy rail 0.15in (+/-0.02) on slides 2-12, absent on 1 and 13.
 6. Zero shapes out of bounds; zero text overflow > 5% (estimate).
 7. Zero empty placeholders.
 8. Titles exact and in order.
 9. Speaker notes >= 20 chars on every slide.
10. Logo on slide 1 (top-left) and slide 13 (bottom-center), aspect within 5%,
    width >= 0.8in.
"""

from __future__ import annotations

import math
import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

EMU_PER_IN = 914400
PT_PER_IN = 72

PALETTE = {"070A40", "E63946", "00BFFF", "F2F2F2", "FFFFFF", "6B7280"}
# The logo PNG has a #02062F background removed via a:clrChange on slides 1/13.
LOGO_BG = "02062F"
ALLOWED = PALETTE | {LOGO_BG}

SLIDE_W_IN = 10.0
SLIDE_H_IN = 7.5
RAIL_W_IN = 0.15

EXPECTED_TITLES = [
    "AI Venture Studio",
    "Executive Summary",
    "The Studio Operating Engine",
    "Portfolio Archetypes A\u2013D",
    "Stage-Gated Development Lifecycle",
    "Technical Architecture (AI-Agnostic System)",
    "Core Performance Metrics",
    "Go-to-Market & Client Acquisition",
    "Governance & Compliance",
    "What You Get as a Studio Partner",
    "Next Steps",
    "Brand Compliance Checklist",
    "Thank You",
]

LOGO_ASPECT = 880 / 713  # cropped content bbox of logo-full.png


def emu_to_in(v: int | None) -> float:
    return (v or 0) / EMU_PER_IN


def shape_text(shape) -> str:
    if not shape.has_text_frame:
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs).strip()


def is_picture(shape) -> bool:
    return shape.shape_type == MSO_SHAPE_TYPE.PICTURE


def iter_slide_shapes(slide):
    """Yield every shape, descending into groups."""
    for shape in slide.shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                yield child


def check_count(prs):
    n = len(prs.slides)
    return n == 13, f"{n} slides (expected 13)"


def check_size(prs):
    w = emu_to_in(prs.slide_width)
    h = emu_to_in(prs.slide_height)
    ok = abs(w - SLIDE_W_IN) < 0.01 and abs(h - SLIDE_H_IN) < 0.01
    return ok, f"{w:.2f} x {h:.2f} in (expected 10.00 x 7.50)"


def check_fonts(prs):
    bad = []
    for idx, slide in enumerate(prs.slides, start=1):
        for shape in iter_slide_shapes(slide):
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if not run.text.strip():
                        continue
                    name = run.font.name
                    if name != "Arial":
                        bad.append((idx, name, run.text[:30]))
    return not bad, f"{len(bad)} non-Arial runs" + (f"; first: {bad[:3]}" if bad else "")


def check_palette(prs):
    bad = []
    for idx, slide in enumerate(prs.slides, start=1):
        for el in slide._element.iter(qn("a:srgbClr")):
            val = (el.get("val") or "").upper()
            if val and val not in ALLOWED:
                bad.append((idx, val))
    return not bad, f"{len(bad)} off-palette RGB values" + (f"; {bad[:5]}" if bad else "")


def find_rail(slide):
    for shape in iter_slide_shapes(slide):
        if not shape.has_text_frame:
            continue
        x = emu_to_in(shape.left)
        w, h = emu_to_in(shape.width), emu_to_in(shape.height)
        if abs(x) < 0.01 and abs(w - RAIL_W_IN) <= 0.02 and abs(h - SLIDE_H_IN) < 0.02:
            return shape
    return None


def check_rail(prs):
    problems = []
    for idx, slide in enumerate(prs.slides, start=1):
        rail = find_rail(slide)
        if 2 <= idx <= 12:
            if rail is None:
                problems.append(f"s{idx}: rail missing")
            else:
                w = emu_to_in(rail.width)
                if abs(w - RAIL_W_IN) > 0.02:
                    problems.append(f"s{idx}: rail width {w:.3f}in")
        else:
            if rail is not None:
                problems.append(f"s{idx}: rail present")
    return not problems, "rail geometry ok" if not problems else "; ".join(problems)


def estimate_text_height(shape) -> float:
    """Rough wrapped-text height in inches for an Arial text frame."""
    if not shape.has_text_frame:
        return 0.0
    tf = shape.text_frame
    avail_w_pt = emu_to_in(shape.width) * PT_PER_IN
    avail_w_pt -= (tf.margin_left + tf.margin_right) / EMU_PER_IN * PT_PER_IN
    total_pt = 0.0
    for para in tf.paragraphs:
        text = "".join(r.text for r in para.runs)
        if not text.strip():
            continue
        size = 14.0
        for run in para.runs:
            if run.font.size is not None:
                size = run.font.size.pt
                break
        spacing = para.line_spacing if isinstance(para.line_spacing, float) else 1.2
        cpl = max(1, int(avail_w_pt / (0.52 * size)))
        lines = max(1, math.ceil(len(text) / cpl))
        total_pt += lines * size * spacing
        if para.space_after is not None:
            total_pt += para.space_after.pt
    return total_pt / PT_PER_IN


def check_bounds_and_overflow(prs):
    oob = []
    overflow = []
    for idx, slide in enumerate(prs.slides, start=1):
        for shape in iter_slide_shapes(slide):
            x, y = emu_to_in(shape.left), emu_to_in(shape.top)
            w, h = emu_to_in(shape.width), emu_to_in(shape.height)
            if x < -0.01 or y < -0.01 or x + w > SLIDE_W_IN + 0.01 or y + h > SLIDE_H_IN + 0.01:
                oob.append(f"s{idx} {shape.shape_type} ({x:.2f},{y:.2f},{w:.2f},{h:.2f})")
            if shape.has_text_frame and h > 0:
                need = estimate_text_height(shape)
                if need > h * 1.05:
                    label = shape_text(shape).replace("\n", " ")[:28]
                    overflow.append(
                        f"s{idx} '{label}' needs {need:.2f}in in {h:.2f}in ({need / h:.0%})"
                    )
    ok = not oob and not overflow
    detail = []
    if oob:
        detail.append(f"{len(oob)} out-of-bounds: {oob[:3]}")
    if overflow:
        detail.append(f"{len(overflow)} overflow: {overflow[:3]}")
    return ok, "; ".join(detail) if detail else "bounds and text fit ok"


def check_placeholders(prs):
    empty = []
    for idx, slide in enumerate(prs.slides, start=1):
        for shape in iter_slide_shapes(slide):
            if shape.is_placeholder and not shape_text(shape):
                empty.append(f"s{idx} {shape.placeholder_format.type}")
    return not empty, f"{len(empty)} empty placeholders" + (f"; {empty}" if empty else "")


def check_titles(prs):
    found = []
    for idx, slide in enumerate(prs.slides, start=1):
        shapes = list(iter_slide_shapes(slide))
        if idx in (1, len(EXPECTED_TITLES)):
            shapes = [sh for sh in shapes if sh.has_text_frame]
        else:
            shapes = [sh for sh in shapes if sh.has_text_frame and emu_to_in(sh.top) < 1.15]
        found.append([shape_text(sh) for sh in shapes])
    problems = []
    for i, expected in enumerate(EXPECTED_TITLES):
        if i >= len(found):
            problems.append(f"s{i + 1}: slide missing")
            continue
        if expected not in found[i]:
            problems.append(f"s{i + 1}: expected {expected!r}")
    return not problems, "titles in order" if not problems else "; ".join(problems)


def check_notes(prs):
    bad = []
    for idx, slide in enumerate(prs.slides, start=1):
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        if len(notes) < 20:
            bad.append(f"s{idx}: {len(notes)} chars")
    return not bad, "notes present on all slides" if not bad else "; ".join(bad)


def check_logo(prs):
    slides = list(prs.slides)
    problems = []

    def pic_info(shape):
        w, h = emu_to_in(shape.width), emu_to_in(shape.height)
        left, top = emu_to_in(shape.left), emu_to_in(shape.top)
        crop_l = shape.crop_left or 0.0
        crop_r = shape.crop_right or 0.0
        crop_t = shape.crop_top or 0.0
        crop_b = shape.crop_bottom or 0.0
        vis_w = w * (1.0 - crop_l - crop_r)
        vis_h = h * (1.0 - crop_t - crop_b)
        vis_left = left + crop_l * w
        vis_top = top + crop_t * h
        return vis_left, vis_top, vis_w, vis_h

    pics1 = [pic_info(sh) for sh in iter_slide_shapes(slides[0]) if is_picture(sh)]
    if not pics1:
        problems.append("s1: no logo picture")
    else:
        x, y, w, h = pics1[0]
        if not (x < 1.0 and y < 1.2):
            problems.append(f"s1: logo not top-left ({x:.2f},{y:.2f})")
        if w < 0.8:
            problems.append(f"s1: logo width {w:.2f}in")
        if abs((w / h) / LOGO_ASPECT - 1) > 0.05:
            problems.append(f"s1: aspect {w / h:.3f} vs {LOGO_ASPECT:.3f}")

    pics13 = [pic_info(sh) for sh in iter_slide_shapes(slides[12]) if is_picture(sh)]
    if not pics13:
        problems.append("s13: no logo picture")
    else:
        x, y, w, h = pics13[0]
        if abs((x + w / 2) - SLIDE_W_IN / 2) > 0.1:
            problems.append(f"s13: logo not horizontally centered (cx={x + w / 2:.2f})")
        if y < SLIDE_H_IN * 0.5:
            problems.append(f"s13: logo not in lower half (y={y:.2f})")
        if w < 0.8:
            problems.append(f"s13: logo width {w:.2f}in")
        if abs((w / h) / LOGO_ASPECT - 1) > 0.05:
            problems.append(f"s13: aspect {w / h:.3f} vs {LOGO_ASPECT:.3f}")

    return not problems, "logo placement ok" if not problems else "; ".join(problems)


def main() -> int:
    path = sys.argv[1] if len(sys.argv) > 1 else "ai_venture_studio_client_presentation.pptx"
    if not os.path.exists(path):
        print(f"ERROR: deck not found: {path}")
        return 1
    prs = Presentation(path)

    checks = [
        ("13 slides", check_count(prs)),
        ("10 x 7.5 in", check_size(prs)),
        ("Arial only", check_fonts(prs)),
        ("palette only", check_palette(prs)),
        ("navy rail", check_rail(prs)),
        ("bounds / overflow", check_bounds_and_overflow(prs)),
        ("no empty placeholders", check_placeholders(prs)),
        ("title sequence", check_titles(prs)),
        ("speaker notes", check_notes(prs)),
        ("logo placement", check_logo(prs)),
    ]

    print(f"=== QA: {os.path.basename(path)} ===")
    failed = 0
    for name, (ok, detail) in checks:
        print(f"[{'PASS' if ok else 'FAIL'}] {name}: {detail}")
        if not ok:
            failed += 1
    print()
    if failed:
        print(f">> {failed}/{len(checks)} CHECKS FAILED <<")
        return 1
    print(f">> ALL {len(checks)} CHECKS PASSED <<")
    return 0


if __name__ == "__main__":
    sys.exit(main())
