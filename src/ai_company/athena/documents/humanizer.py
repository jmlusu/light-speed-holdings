"""
Dehumanization layer for Athena — rewrites AI-generated content to sound human.

This module takes AI-generated resume/cover letter content and rewrites it to
sound natural and human-written while preserving all factual information.
"""

import logging
import re
from dataclasses import dataclass
from enum import Enum
from pathlib import Path
from typing import TYPE_CHECKING, Any, Optional

if TYPE_CHECKING:
    from docx.document import Document as DocxDocumentType

from docx import Document as DocxDocumentClass
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from ai_company.llm.client import LLMClient
from ai_company.model_router import ModelRouter

logger = logging.getLogger(__name__)


class HumanizationIntensity(Enum):
    """Intensity levels for humanization."""

    LIGHT = "light"  # Minor cleanup, preserve structure
    MEDIUM = "medium"  # Balance of naturalness and structure
    HEAVY = "heavy"  # Maximum naturalness, restructure freely


@dataclass
class HumanizationResult:
    """Result of humanization process."""

    original_text: str
    humanized_text: str
    docx_path: Optional[str] = None
    intensity: HumanizationIntensity = HumanizationIntensity.MEDIUM
    changes_made: Optional[list[str]] = None

    def __post_init__(self) -> None:
        if self.changes_made is None:
            self.changes_made = []


# ─── System Prompts ────────────────────────────────────────────────

SYSTEM_PROMPT_BASE = """You are an expert resume and cover letter writer who specializes in making AI-generated content sound authentically human. Your goal is to rewrite the provided text so it reads like a real person wrote it — confident, natural, and professional — while preserving EVERY SINGLE FACT, DATE, NUMBER, COMPANY NAME, AND METRIC exactly as provided.

CORE PRINCIPLES:
1. PRESERVE ALL FACTS: Dates, company names, titles, metrics, percentages, dollar amounts, team sizes, project names — NOTHING CHANGES.
2. SOUND HUMAN: Use varied sentence structures. Mix short and long sentences. Use natural transitions. Avoid robotic patterns.
3. REMOVE AI TELLS: Strip out corporate speak, buzzwords, clichés, and telltale AI patterns.
4. STAY PROFESSIONAL: Confident but not arrogant. Authentic but polished.
5. KEEP STRUCTURE: Maintain the document's logical flow (intro → body → close for cover letters; sections for resumes).

SPECIFIC AI PATTERNS TO ELIMINATE:
- "I am writing to express my interest in..." → Natural, specific opening
- "I have extensive experience in..." → Concrete examples with numbers
- "Passionate about..." / "Excited to..." → Show, don't tell
- "Leveraged", "spearheaded", "orchestrated", "delivered", "drove", "championed" — vary your verbs
- Bullet points all starting with the same action verb
- "In today's fast-paced world..." / "In the ever-evolving landscape..." — delete entirely
- "As a [title], I..." — vary sentence openings
- "I am confident that..." / "I believe that..." — state it directly
- Overuse of "proven track record", "results-oriented", "dynamic", "strategic"
- Generic filler: "utilize" → "use", "facilitate" → "help", "implement" → "build/do"
- "I would welcome the opportunity to discuss..." → "I'd love to talk about..."

NATURAL REPLACEMENTS:
- Instead of "spearheaded" every bullet: "led", "built", "created", "launched", "ran", "owned", "directed", "managed", "designed", "developed"
- Instead of "I have X years of experience in Y": "I've spent X years doing Y at [company], where I [specific thing]"
- Instead of "passionate about": "I've worked on", "I've built", "I care about", "what draws me to"
- Vary sentence starts: "At [company], I...", "My work on [project]...", "When I [did X]...", "The [project] I led..."

TONE: Write like a senior professional talking to a peer over coffee — competent, specific, not trying to impress with vocabulary."""


SYSTEM_PROMPT_LIGHT = (
    SYSTEM_PROMPT_BASE
    + """

INTENSITY: LIGHT
- Fix only the most obvious AI tells
- Keep original structure and paragraph breaks
- Minimal rewriting — just smooth the rough edges
- Target: Sounds like a human edited an AI draft"""
)


SYSTEM_PROMPT_MEDIUM = (
    SYSTEM_PROMPT_BASE
    + """

INTENSITY: MEDIUM
- Rewrite sentences for natural flow and varied structure
- Reorganize bullet points to avoid repetitive patterns
- Replace buzzwords with concrete language
- Add natural transitions between paragraphs
- Target: Sounds like a human wrote it from scratch"""
)


SYSTEM_PROMPT_HEAVY = (
    SYSTEM_PROMPT_BASE
    + """

INTENSITY: HEAVY
- Full rewrite with maximum naturalness
- Restructure paragraphs for better flow
- Inject personality and voice
- Vary sentence length dramatically (short punchy sentences mixed with longer ones)
- Use contractions naturally (I've, I'm, it's, don't)
- Add subtle rhetorical devices (parallelism, strategic fragments)
- Target: Indistinguishable from a strong human writer"""
)


# ─── User Prompt Templates ────────────────────────────────────────

COVER_LETTER_USER_PROMPT = """Rewrite this cover letter to sound human.

ORIGINAL CONTENT:
{content}

TARGET JOB DESCRIPTION:
{job_description}

USER PROFILE:
- Name: {full_name}
- Headline: {headline}
- Key Skills: {skills}
- Years Experience: {years_experience}
- Top Achievements: {achievements}

REQUIREMENTS:
- Keep ALL facts, dates, companies, metrics exactly the same
- Match the job's tone and keywords naturally
- Sound like {full_name} actually wrote this
- Professional but not stiff
- {intensity_instruction}"""

RESUME_SECTION_USER_PROMPT = """Rewrite this resume section to sound human.

ORIGINAL CONTENT:
{content}

SECTION TYPE: {section_type}

USER PROFILE:
- Name: {full_name}
- Headline: {headline}
- Key Skills: {skills}
- Years Experience: {years_experience}

REQUIREMENTS:
- Keep ALL facts, dates, companies, metrics, titles exactly the same
- For experience: vary action verbs across bullets, add specific outcomes
- For skills: group naturally, don't just list
- For summary: make it sound like a real person's voice
- {intensity_instruction}"""


GENERIC_USER_PROMPT = """Rewrite this content to sound human.

ORIGINAL CONTENT:
{content}

CONTEXT: {context}

USER PROFILE:
- Name: {full_name}
- Headline: {headline}

REQUIREMENTS:
- Keep ALL facts, dates, numbers, names exactly the same
- {intensity_instruction}"""


INTENSITY_INSTRUCTIONS = {
    HumanizationIntensity.LIGHT: "Make minimal changes — just fix the most obvious AI patterns and awkward phrasing.",
    HumanizationIntensity.MEDIUM: "Rewrite for natural flow. Vary sentence structure, replace buzzwords, improve transitions.",
    HumanizationIntensity.HEAVY: "Full rewrite. Restructure freely for maximum naturalness. Use contractions, varied sentence lengths, and authentic voice.",
}


# ─── AI Pattern Detection (for analysis) ──────────────────────────

AI_TELL_PATTERNS = {
    "generic_opening": [
        r"I am writing to express my interest",
        r"I am writing to apply for",
        r"I am excited to apply",
        r"I am submitting my application",
    ],
    "vague_experience": [
        r"I have extensive experience in",
        r"I have a proven track record of",
        r"I possess strong experience in",
        r"My background includes extensive",
    ],
    "passion_cliche": [
        r"passionate about",
        r"deeply passionate about",
        r"excited about the opportunity",
        r"eager to bring my passion",
    ],
    "overused_verbs": [
        r"\bspearheaded\b",
        r"\borchestrated\b",
        r"\bleveraged\b",
        r"\bdelivered\b",
        r"\bdrove\b",
        r"\bchampioned\b",
        r"\boptimized\b",
        r"\btransformed\b",
        r"\brevolutionized\b",
    ],
    "corporate_speak": [
        r"\butilize\b",
        r"\bfacilitate\b",
        r"\bimplement\b",
        r"\bstrategic initiative\b",
        r"\bcross-functional\b",
        r"\bstakeholder\b",
        r"\bholistic approach\b",
        r"\bend-to-end\b",
        r"\bdeep dive\b",
        r"\bmoving forward\b",
    ],
    "filler_phrases": [
        r"in today's (fast-paced|ever-evolving|rapidly changing)",
        r"in the (ever-evolving|rapidly changing|dynamic) landscape",
        r"as a .+, I ",
        r"I am confident that",
        r"I believe that I",
        r"I would welcome the opportunity",
        r"thank you for your consideration",
        r"please do not hesitate",
    ],
    "weak_closings": [
        r"I look forward to hearing from you",
        r"I hope to discuss this further",
        r"thank you for your time and consideration",
    ],
}


# ─── Humanizer Class ──────────────────────────────────────────────


class Humanizer:
    """
    Rewrites AI-generated resume/cover letter content to sound human.

    Uses LLMClient for rewriting with intensity-controlled prompts.
    Preserves all factual information while removing AI tells.
    """

    def __init__(
        self,
        config_path: str = "company/models.yaml",
        registry_path: str = "company/agent-registry.json",
    ) -> None:
        self.llm_client = LLMClient(config_path=config_path, registry_path=registry_path)
        self.model_router = ModelRouter(config_path=config_path, registry_path=registry_path)

    def humanize_cover_letter(
        self,
        content: str,
        job_description: str,
        user_profile: dict[str, Any],
        intensity: HumanizationIntensity = HumanizationIntensity.MEDIUM,
        agent_name: str = "athena-document-writer",
    ) -> HumanizationResult:
        """Humanize a cover letter."""
        system_prompt = self._get_system_prompt(intensity)
        intensity_instruction = INTENSITY_INSTRUCTIONS[intensity]

        user_prompt = COVER_LETTER_USER_PROMPT.format(
            content=content,
            job_description=job_description[:3000],  # Truncate if very long
            full_name=user_profile.get("full_name", "the applicant"),
            headline=user_profile.get("headline", ""),
            skills=", ".join(user_profile.get("skills", [])[:15]),
            years_experience=user_profile.get("years_experience", "several"),
            achievements="; ".join(user_profile.get("achievements", [])[:5]),
            intensity_instruction=intensity_instruction,
        )

        humanized = self._call_llm(
            agent_name=agent_name,
            system_prompt=system_prompt,
            user_prompt=user_prompt,
        )

        changes = self._detect_changes(content, humanized)
        return HumanizationResult(
            original_text=content,
            humanized_text=humanized,
            intensity=intensity,
            changes_made=changes,
        )

    def humanize_resume_section(
        self,
        content: str,
        section_type: str,  # "experience", "summary", "skills", "education", etc.
        user_profile: dict[str, Any],
        intensity: HumanizationIntensity = HumanizationIntensity.MEDIUM,
        agent_name: str = "athena-document-writer",
    ) -> HumanizationResult:
        """Humanize a resume section."""
        system_prompt = self._get_system_prompt(intensity)
        intensity_instruction = INTENSITY_INSTRUCTIONS[intensity]

        user_prompt = RESUME_SECTION_USER_PROMPT.format(
            content=content,
            section_type=section_type,
            full_name=user_profile.get("full_name", "the applicant"),
            headline=user_profile.get("headline", ""),
            skills=", ".join(user_profile.get("skills", [])[:15]),
            years_experience=user_profile.get("years_experience", "several"),
            intensity_instruction=intensity_instruction,
        )

        humanized = self._call_llm(
            agent_name=agent_name,
            system_prompt=system_prompt,
            user_prompt=user_prompt,
        )

        changes = self._detect_changes(content, humanized)
        return HumanizationResult(
            original_text=content,
            humanized_text=humanized,
            intensity=intensity,
            changes_made=changes,
        )

    def humanize_generic(
        self,
        content: str,
        context: str,
        user_profile: dict[str, Any],
        intensity: HumanizationIntensity = HumanizationIntensity.MEDIUM,
        agent_name: str = "athena-document-writer",
    ) -> HumanizationResult:
        """Humanize any text content."""
        system_prompt = self._get_system_prompt(intensity)
        intensity_instruction = INTENSITY_INSTRUCTIONS[intensity]

        user_prompt = GENERIC_USER_PROMPT.format(
            content=content,
            context=context,
            full_name=user_profile.get("full_name", "the applicant"),
            headline=user_profile.get("headline", ""),
            intensity_instruction=intensity_instruction,
        )

        humanized = self._call_llm(
            agent_name=agent_name,
            system_prompt=system_prompt,
            user_prompt=user_prompt,
        )

        changes = self._detect_changes(content, humanized)
        return HumanizationResult(
            original_text=content,
            humanized_text=humanized,
            intensity=intensity,
            changes_made=changes,
        )

    def _get_system_prompt(self, intensity: HumanizationIntensity) -> str:
        """Get the appropriate system prompt for intensity."""
        return {
            HumanizationIntensity.LIGHT: SYSTEM_PROMPT_LIGHT,
            HumanizationIntensity.MEDIUM: SYSTEM_PROMPT_MEDIUM,
            HumanizationIntensity.HEAVY: SYSTEM_PROMPT_HEAVY,
        }[intensity]

    def _call_llm(
        self,
        agent_name: str,
        system_prompt: str,
        user_prompt: str,
    ) -> str:
        """Call the LLM and return the humanized text."""
        # For humanization, we want the raw text response, not JSON
        # Use the LLMClient's provider directly with a simple chat call
        route = self.model_router.resolve(
            agent_name=agent_name,
            priority="medium",
            context="domain_code_review",  # Use premium tier for quality writing
            task_prompt=user_prompt,
        )

        provider = self.llm_client.get_provider(route.provider)
        if not provider:
            raise RuntimeError(f"Provider {route.provider} not available")

        logger.info(
            "Humanizing with %s/%s (tier: %s)",
            route.provider,
            route.model,
            route.tier,
        )

        response = provider.chat(
            system_prompt=system_prompt,
            user_prompt=user_prompt,
            model=route.model,
        )

        return response.content.strip()

    def _detect_changes(self, original: str, humanized: str) -> list[str]:
        """Detect what changes were made."""
        changes = []

        # Check for AI tell removal
        for category, patterns in AI_TELL_PATTERNS.items():
            for pattern in patterns:
                if re.search(pattern, original, re.IGNORECASE) and not re.search(
                    pattern, humanized, re.IGNORECASE
                ):
                    changes.append(f"Removed {category}: '{pattern}'")

        # Check sentence structure variation
        orig_sentences = re.split(r"[.!?]+", original)
        hum_sentences = re.split(r"[.!?]+", humanized)
        if len(orig_sentences) != len(hum_sentences):
            changes.append(f"Sentence count changed: {len(orig_sentences)} → {len(hum_sentences)}")

        # Check for contractions (sign of naturalness)
        contractions = ["I've", "I'm", "it's", "don't", "won't", "can't", "I'll", "you're", "we've"]
        orig_contractions = sum(1 for c in contractions if c in original)
        hum_contractions = sum(1 for c in contractions if c in humanized)
        if hum_contractions > orig_contractions:
            changes.append(f"Added natural contractions: {orig_contractions} → {hum_contractions}")

        # Check verb diversity in bullet points
        orig_bullets = re.findall(r"^[\s•\-]\s*(\w+)", original, re.MULTILINE)
        hum_bullets = re.findall(r"^[\s•\-]\s*(\w+)", humanized, re.MULTILINE)
        if orig_bullets and hum_bullets:
            orig_unique = len(set(v.lower() for v in orig_bullets))
            hum_unique = len(set(v.lower() for v in hum_bullets))
            if hum_unique > orig_unique:
                changes.append(
                    f"Increased bullet verb diversity: {orig_unique} → {hum_unique} unique verbs"
                )

        return changes

    def to_docx(
        self,
        result: HumanizationResult,
        output_path: str | Path,
        document_type: str = "cover_letter",  # or "resume"
        title: str = "",
    ) -> str:
        """Convert humanized text to a formatted DOCX file."""
        doc = DocxDocumentClass()
        path = Path(output_path)

        # Set default style
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15

        if document_type == "cover_letter":
            self._format_cover_letter(doc, result.humanized_text, title)
        else:
            self._format_resume(doc, result.humanized_text, title)

        doc.save(str(path))
        return str(path)

    def _format_cover_letter(self, doc: "DocxDocumentType", text: str, title: str) -> None:
        """Format text as a cover letter."""
        # Title
        if title:
            heading = doc.add_heading(title, level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in heading.runs:
                run.font.size = Pt(14)
                run.font.color.rgb = None  # Use default

        # Split into paragraphs and add
        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]

        for i, para_text in enumerate(paragraphs):
            # Handle bullet points
            if para_text.startswith(("•", "-", "*")):
                for line in para_text.split("\n"):
                    line = line.strip()
                    if line:
                        p = doc.add_paragraph(style="List Bullet")
                        p.add_run(line.lstrip("•-* ").strip())
            else:
                p = doc.add_paragraph(para_text)
                p.paragraph_format.first_line_indent = Inches(0.5) if i > 0 else Inches(0)

    def _format_resume(self, doc: "DocxDocumentType", text: str, title: str) -> None:
        """Format text as a resume."""
        if title:
            heading = doc.add_heading(title, level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in heading.runs:
                run.font.size = Pt(16)
                run.bold = True

        # Simple paragraph-based formatting for resume sections
        sections = text.split("\n\n")
        for section in sections:
            section = section.strip()
            if not section:
                continue

            lines = section.split("\n")
            # First line might be a section header
            if len(lines) > 1 and len(lines[0]) < 80 and lines[0].isupper():
                heading = doc.add_heading(lines[0], level=2)
                for run in heading.runs:
                    run.font.size = Pt(12)
                for line in lines[1:]:
                    if line.strip().startswith(("•", "-", "*")):
                        p = doc.add_paragraph(style="List Bullet")
                        p.add_run(line.lstrip("•-* ").strip())
                    elif line.strip():
                        doc.add_paragraph(line.strip())
            else:
                for line in lines:
                    if line.strip().startswith(("•", "-", "*")):
                        p = doc.add_paragraph(style="List Bullet")
                        p.add_run(line.lstrip("•-* ").strip())
                    elif line.strip():
                        doc.add_paragraph(line.strip())

    def analyze_ai_tells(self, text: str) -> dict[str, list[str]]:
        """Analyze text for AI tells and return findings by category."""
        findings: dict[str, list[str]] = {}

        for category, patterns in AI_TELL_PATTERNS.items():
            matches = []
            for pattern in patterns:
                for match in re.finditer(pattern, text, re.IGNORECASE):
                    # Get context around match
                    start = max(0, match.start() - 50)
                    end = min(len(text), match.end() + 50)
                    context = text[start:end].replace("\n", " ")
                    matches.append(f"...{context}...")
            if matches:
                findings[category] = matches

        return findings


# ─── Convenience Functions ────────────────────────────────────────


async def humanize_cover_letter_async(
    content: str,
    job_description: str,
    user_profile: dict[str, Any],
    intensity: HumanizationIntensity = HumanizationIntensity.MEDIUM,
    output_path: Optional[str] = None,
) -> HumanizationResult:
    """Async wrapper for humanizing cover letters (if needed)."""
    # The LLMClient is sync, but we can wrap for async contexts
    humanizer = Humanizer()
    result = humanizer.humanize_cover_letter(
        content=content,
        job_description=job_description,
        user_profile=user_profile,
        intensity=intensity,
    )

    if output_path:
        result.docx_path = humanizer.to_docx(result, output_path, "cover_letter")

    return result


def create_humanizer(
    config_path: str = "company/models.yaml",
    registry_path: str = "company/agent-registry.json",
) -> Humanizer:
    """Factory function to create a Humanizer instance."""
    return Humanizer(config_path=config_path, registry_path=registry_path)


# ─── Export ────────────────────────────────────────────────────────

__all__ = [
    "Humanizer",
    "HumanizationIntensity",
    "HumanizationResult",
    "create_humanizer",
    "humanize_cover_letter_async",
]
