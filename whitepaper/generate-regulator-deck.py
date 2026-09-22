"""
LightSpeed Holdings — Malawi Regulator Deck Generator
Creates a 14-slide regulator briefing focused on LightSpeed's governed, AI-native
potential in Malawi for MACRA / digital & data protection regulators.

Open by design: this deck is generated from a committed, readable script
(no binary-only assets; brand tokens are the single source of truth).

Brand: LightSpeed Holdings Limited | Tagline: Aspire. Act. Achieve.
Colors: Navy #070A40, Red #E63946, Cyan #00BFFF, Grey #F2F2F2
"""

import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# Brand Colors
NAVY = RGBColor(0x07, 0x0A, 0x40)
RED = RGBColor(0xE6, 0x39, 0x46)
CYAN = RGBColor(0x00, 0xBF, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0xF2, 0xF2, 0xF2)
DARK_GREY = RGBColor(0x6B, 0x72, 0x80)
LIGHT_TEXT = RGBColor(0x9C, 0xA3, 0xAF)


def set_slide_bg(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, shape_type, left, top, width, height, fill_color=None, line_color=None):
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def add_text_box(slide, left, top, width, height, text, font_size=14, bold=False,
                 color=NAVY, alignment=PP_ALIGN.LEFT, font_name="Arial"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_list(slide, left, top, width, height, items, font_size=12, color=NAVY):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Arial"
        p.space_after = Pt(6)
        p.level = 0
    return txBox


def add_kpi_card(slide, x, y, value, label):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.8), Inches(1.3), fill_color=GREY)
    add_text_box(slide, x + Inches(0.15), y + Inches(0.15), Inches(2.5), Inches(0.6),
                 value, font_size=26, bold=True, color=NAVY, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.15), y + Inches(0.75), Inches(2.5), Inches(0.4),
                 label, font_size=11, color=DARK_GREY, alignment=PP_ALIGN.CENTER)


def add_callout(slide, top, text, height=1.5):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(top), Inches(8.0), Inches(height),
              fill_color=GREY)
    add_text_box(slide, Inches(1.3), Inches(top + 0.2), Inches(7.4), Inches(height - 0.3),
                 text, font_size=13, bold=True, color=NAVY, alignment=PP_ALIGN.CENTER)


def add_content_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.15), Inches(7.5), fill_color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(0.4), Inches(9), Inches(0.6),
                 title, font_size=28, bold=True, color=NAVY)
    return slide


# ── SLIDE 1: TITLE ──────────────────────────────────────────
def create_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, NAVY)

    add_text_box(slide, Inches(1), Inches(1.4), Inches(8), Inches(0.6),
                 "LIGHTSPEED HOLDINGS LIMITED\u2122", font_size=32, bold=True, color=WHITE,
                 alignment=PP_ALIGN.CENTER)

    add_text_box(slide, Inches(1), Inches(2.1), Inches(8), Inches(0.5),
                 "Aspire. Act. Achieve.", font_size=18, color=CYAN,
                 alignment=PP_ALIGN.CENTER)

    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(3.5), Inches(2.8), Inches(3), Inches(0.04),
              fill_color=RED)

    add_text_box(slide, Inches(1), Inches(3.2), Inches(8), Inches(1),
                 "Regulator Briefing\nMalawi\u2019s Governed AI-Native Potential",
                 font_size=24, bold=False, color=WHITE, alignment=PP_ALIGN.CENTER)

    add_text_box(slide, Inches(1), Inches(4.9), Inches(8), Inches(0.5),
                 "Open by design. Governed by construction.",
                 font_size=14, color=CYAN, alignment=PP_ALIGN.CENTER)

    add_text_box(slide, Inches(1), Inches(5.8), Inches(8), Inches(0.5),
                 "September 2026  |  Prepared for Malawi regulators (MACRA / DPA)",
                 font_size=12, color=LIGHT_TEXT, alignment=PP_ALIGN.CENTER)


# ── SLIDE 2: PURPOSE ────────────────────────────────────────
def create_purpose_slide(prs):
    slide = add_content_slide(prs, "PURPOSE OF THIS BRIEFING")

    items = [
        "Who we are  \u2014  an AI-native services studio incorporated in Malawi",
        "What we operate  \u2014  150+ AI agents across 20 departments, one human CEO",
        "How we are governed  \u2014  human-in-the-loop at every high-stakes step",
        "What compliance means to us  \u2014  Data Protection Act (2024) + GDPR from day one",
        "Why the regulator matters  \u2014  we want to be tested, then held to account",
    ]
    add_bullet_list(slide, Inches(0.8), Inches(1.3), Inches(8.5), Inches(3.5), items, font_size=14, color=NAVY)

    add_callout(slide, 5.1, "This briefing is a transparency artifact.\n"
                            "We publish how we work \u2014 and invite scrutiny.", height=1.5)


# ── SLIDE 3: POLICY WINDOW ──────────────────────────────────
def create_policy_slide(prs):
    slide = add_content_slide(prs, "WHY NOW \u2014 MALAWI\u2019S POLICY WINDOW")

    items = [
        "Data Protection Act (2024) in force since 3 June 2024; MACRA is the designated Data Protection Authority",
        "Legally binding rights already apply to AI: no solely automated decisions with significant effects; DPIA duties for high-risk processing",
        "No dedicated AI statute yet \u2014 the rules are being written now",
        "Draft National AI Strategy + Digital Transformation Strategy under validation (Ministry of Information and Digitalisation, with UNDP)",
        "UNESCO AI Readiness Assessment (RAM) validated July 2026 \u2014 Malawi among six southern African pilot countries",
        "Draft AI Bill announced August 2026 \u2014 to be tabled before December 2026",
        "MACRA published a draft AI regulatory framework for public comment (Sept 2026) \u2014 agentic transparency & consumer protection",
        "Regional anchors: AU Continental AI Strategy (2024); SADC agentic-AI governance framework in drafting",
    ]
    add_bullet_list(slide, Inches(0.8), Inches(1.2), Inches(8.5), Inches(4.6), items, font_size=12, color=NAVY)

    add_callout(slide, 5.9, "LightSpeed wants to co-build this regime \u2014 not wait for it.", height=1.2)


# ── SLIDE 4: WHO WE ARE ─────────────────────────────────────
def create_who_we_are_slide(prs):
    slide = add_content_slide(prs, "WHO WE ARE \u2014 THE AI-NATIVE STUDIO")

    add_text_box(slide, Inches(0.8), Inches(1.15), Inches(8.5), Inches(0.5),
                 "One human CEO owns the outcome. 150+ AI agents execute the work.",
                 font_size=15, bold=False, color=NAVY)

    items = [
        "Delivery pipeline: client brief \u2192 task queue \u2192 assigned agents \u2192 human CEO review \u2192 client deliverable",
        "Departments span engineering, design, marketing, sales, legal, compliance, HR, finance",
        "Every agent action is logged and queryable \u2014 nothing executes silently",
        "We win on speed, cost, and 24/7 capacity \u2014 never by impersonating human staff",
        "Established track record (We Lead Out \u2014 5.0 AppExchange rating; 20+ projects shipped)",
    ]
    add_bullet_list(slide, Inches(0.8), Inches(1.8), Inches(8.5), Inches(3.4), items, font_size=13, color=NAVY)

    add_kpi_card(slide, Inches(0.8), Inches(5.4), "150+", "AI Agents")
    add_kpi_card(slide, Inches(3.9), Inches(5.4), "20", "Departments")
    add_kpi_card(slide, Inches(7.0), Inches(5.4), "1", "Human CEO")


# ── SLIDE 5: DELIVERY MODEL ─────────────────────────────────
def create_delivery_slide(prs):
    slide = add_content_slide(prs, "HOW WE DELIVER \u2014 THE FLOW")

    steps = [
        "1.  Client brief \u2014 scope, price, contract (50% upfront / 50% on delivery)",
        "2.  4-gate onboarding \u2014 contract, data-processing agreement, compliance risk assessment, security review",
        "3.  Task dispatch \u2014 work routed to the right department queue",
        "4.  Agent execution \u2014 audited, logged, bounded by policy and risk tier",
        "5.  Human review \u2014 CEO / executive sign-off on every deliverable",
        "6.  Handover \u2014 client receives a usable artifact + its audit trail",
    ]
    add_bullet_list(slide, Inches(0.8), Inches(1.3), Inches(8.5), Inches(3.8), steps, font_size=14, color=NAVY)

    add_callout(slide, 5.3, "Humans approve. Agents execute. The audit trail proves both.",
                height=1.5)


# ── SLIDE 6: OFFERS ─────────────────────────────────────────
def add_offer_card(slide, x, y, title, price, body):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.9), Inches(2.6), fill_color=GREY)
    add_text_box(slide, x + Inches(0.15), y + Inches(0.1), Inches(2.6), Inches(0.4),
                 title, font_size=12, bold=True, color=RED)
    add_text_box(slide, x + Inches(0.15), y + Inches(0.5), Inches(2.6), Inches(0.3),
                 price, font_size=10, color=NAVY)
    add_text_box(slide, x + Inches(0.15), y + Inches(0.9), Inches(2.6), Inches(1.6),
                 body, font_size=10, color=DARK_GREY)


def create_offers_slide(prs):
    slide = add_content_slide(prs, "WHAT WE DELIVER \u2014 FIVE OFFERS")

    add_offer_card(slide, Inches(0.5), Inches(1.2),
                   "Offer A: Digital Presence", "From MWK 150,000 (~$85)",
                   "Websites, e-commerce, brand identity, Google Business, social media kit")
    add_offer_card(slide, Inches(3.55), Inches(1.2),
                   "Offer B: Process Automation", "From MWK 900,000 (~$500)",
                   "WhatsApp chatbots, document generators, survey automation, dashboards")
    add_offer_card(slide, Inches(6.6), Inches(1.2),
                   "Offer E: Platform Licensing", "$49\u2013$299/mo (SaaS)",
                   "AI Company Builder for agencies & developers \u2014 licensed, audited use")

    add_offer_card(slide, Inches(0.5), Inches(4.0),
                   "Offer C: Data & Reporting", "From MWK 700,000 (~$400)",
                   "Data cleaning, donor reports, dashboards, survey design \u2014 NGO reporting edge")
    add_offer_card(slide, Inches(3.55), Inches(4.0),
                   "Offer D: Digital Marketing", "From MWK 350,000/mo (~$200/mo)",
                   "Social media management, content packs, Google/Facebook ad campaigns")

    add_text_box(slide, Inches(0.5), Inches(6.8), Inches(9), Inches(0.4),
                 "Dual-currency: MWK for local SMEs  |  USD for NGOs & international clients",
                 font_size=12, bold=True, color=NAVY, alignment=PP_ALIGN.CENTER)


# ── SLIDE 7: POTENTIAL FOR MALAWI ───────────────────────────
def create_potential_slide(prs):
    slide = add_content_slide(prs, "POTENTIAL FOR MALAWI")

    items = [
        "SME digitalization at MWK prices \u2014 websites, e-commerce, chatbots previously out of reach",
        "NGO & donor reporting in days, not weeks \u2014 USD revenue into the local economy",
        "Digital skills: Malawians learn to operate agentic systems, not just consume them",
        "Enterprise-grade output without mass hiring \u2014 a capital-light, job-positive model",
        "A proven Malawi playbook that SADC neighbours (16 states, 300M+ people) can adopt",
        "First customer-facing proof targeted by 90 days \u2014 MWK 5,000,000 revenue target",
    ]
    add_bullet_list(slide, Inches(0.8), Inches(1.3), Inches(8.5), Inches(4.2), items, font_size=14, color=NAVY)

    add_callout(slide, 5.7, "If governed AI-native delivery works in Malawi, it works anywhere.",
                height=1.3)


# ── SLIDE 8: GOVERNANCE ─────────────────────────────────────
def create_governance_slide(prs):
    slide = add_content_slide(prs, "GOVERNED BY CONSTRUCTION")

    items = [
        "5-tier approval matrix \u2014 human sign-off scales with risk; high-stakes = mandatory HITL",
        "Risk-classified agent tiers \u2014 autonomous execution only within defined boundaries",
        "Board of Directors with 7 standing committees; Corporate Constitution sets decision order & escalation SLAs",
        "4-gate client onboarding (G1\u2013G4) \u2014 contract, DPA, compliance risk, security review",
        "Programmatic kill switches, circuit breakers, and rate limiters on agent execution",
        "Separation of duties: proposer \u2260 approver \u2260 executor",
    ]
    add_bullet_list(slide, Inches(0.8), Inches(1.3), Inches(8.5), Inches(4.2), items, font_size=13, color=NAVY)

    add_kpi_card(slide, Inches(0.8), Inches(5.6), "5", "Approval Tiers")
    add_kpi_card(slide, Inches(3.9), Inches(5.6), "100%", "High-Stakes Human Review")
    add_kpi_card(slide, Inches(7.0), Inches(5.6), "4", "Onboarding Gates")


# ── SLIDE 9: DATA PROTECTION ────────────────────────────────
def create_data_protection_slide(prs):
    slide = add_content_slide(prs, "DATA PROTECTION & PRIVACY BY DESIGN")

    items = [
        "Compliant with the Data Protection Act (2024) and GDPR from day one \u2014 not retrofitted",
        "Client data is never used to train models",
        "Sovereign in-country processing; offline-first architecture with local models where required",
        "Cross-border LLM flows documented, risk-assessed, and consent-based",
        "High-risk processing subject to Data Protection Impact Assessment (DPIA)",
        "Respects the right not to be subject to solely automated decisions with significant effects",
    ]
    add_bullet_list(slide, Inches(0.8), Inches(1.3), Inches(8.5), Inches(4.2), items, font_size=13, color=NAVY)

    add_callout(slide, 5.7, "We treat the DPA\u2019s automated-decision and DPIA duties as the "
                            "specification for agentic systems.", height=1.3)


# ── SLIDE 10: TRANSPARENCY & CONSUMER PROTECTION ────────────
def create_transparency_slide(prs):
    slide = add_content_slide(prs, "TRANSPARENCY & CONSUMER PROTECTION")

    items = [
        "Clients are told when AI is involved \u2014 no impersonation of human staff",
        "Honesty badges on every claim: source-tagged, confidence-labelled output",
        "A human can always be reached \u2014 chat and WhatsApp flows hand off to a person",
        "Clear recourse: every deliverable has an accountable human + a dispute channel",
        "Consent-based data use; customer data serves the service, not ad models",
        "Aligned with MACRA\u2019s draft framework themes: agentic transparency + consumer protection",
    ]
    add_bullet_list(slide, Inches(0.8), Inches(1.3), Inches(8.5), Inches(4.2), items, font_size=13, color=NAVY)

    add_callout(slide, 5.7, "A consumer should never need to guess whether they are talking to a machine \u2014 or to whom they can complain.",
                height=1.3)


# ── SLIDE 11: AUDIT & ACCOUNTABILITY ────────────────────────
def create_audit_slide(prs):
    slide = add_content_slide(prs, "AUDIT & ACCOUNTABILITY")

    items = [
        "Immutable audit trail \u2014 every agent action, decision, and approval logged and queryable",
        "RACI matrices \u2014 who is accountable for every workflow step",
        "Escalation SLAs and a postmortem process for incidents",
        "1,800+ automated tests; ruff + mypy + bandit clean on the platform",
        "Agency-as-instrument: agents hold no legal personality; liability rests with LightSpeed Holdings Limited as the deploying entity",
        "Regulator access: technical read of control architecture, on demand",
    ]
    add_bullet_list(slide, Inches(0.8), Inches(1.3), Inches(8.5), Inches(4.2), items, font_size=13, color=NAVY)

    add_kpi_card(slide, Inches(0.8), Inches(5.6), "1,800+", "Automated Tests")
    add_kpi_card(slide, Inches(3.9), Inches(5.6), "100%", "Actions Logged")
    add_kpi_card(slide, Inches(7.0), Inches(5.6), "1", "Accountable Entity")


# ── SLIDE 12: STRATEGIC ALIGNMENT ───────────────────────────
def create_alignment_slide(prs):
    slide = add_content_slide(prs, "ALIGNED WITH NATIONAL & REGIONAL STRATEGY")

    items = [
        "Malawi National AI Strategy  \u2014  submitted formal consultation comments; proposed governance architecture & lighthouse use cases",
        "UNESCO AI Readiness (RAM)  \u2014  readiness to operate under the validated assessment for southern Africa",
        "SADC Agentic AI Governance Framework  \u2014  graduated autonomy tiers, machine-readable agent credentials, immutable audit trails, data sovereignty",
        "AU Continental AI Strategy (2024)  \u2014  aligned operational posture",
        "Cross-border principle  \u2014  agents hosted in Malawi execute abroad under documented, consent-based flows",
    ]
    add_bullet_list(slide, Inches(0.8), Inches(1.3), Inches(8.5), Inches(4.2), items, font_size=13, color=NAVY)

    add_callout(slide, 5.7, "We do not wait for the rules to be written. We build the controls the rules will require.",
                height=1.3)


# ── SLIDE 13: WORKING WITH THE REGULATOR ────────────────────
def create_regulator_slide(prs):
    slide = add_content_slide(prs, "HOW WE WISH TO WORK WITH MACRA")

    items = [
        "Technical briefings  \u2014  agentic architecture, audit trails, and controls, explained by the operator",
        "Consultation partner  \u2014  substantive input to the National AI Strategy and draft AI framework",
        "Lighthouse pilot  \u2014  a citizen-inquiry or legislative-summary agent for a public-sector partner, fully supervised",
        "Transparency on demand  \u2014  we publish how we work and open the control architecture to review",
        "Report line  \u2014  a named compliance contact for the regulator to reach at any time",
    ]
    add_bullet_list(slide, Inches(0.8), Inches(1.3), Inches(8.5), Inches(4.2), items, font_size=14, color=NAVY)

    add_callout(slide, 5.7, "The ask is simple: engage with us, test us, then hold us to the standard we publish.",
                height=1.3)


# ── SLIDE 14: THANK YOU ─────────────────────────────────────
def create_thank_you_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, NAVY)

    add_text_box(slide, Inches(1), Inches(1.9), Inches(8), Inches(1),
                 "THANK YOU", font_size=36, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)

    add_text_box(slide, Inches(1), Inches(3.1), Inches(8), Inches(0.6),
                 "Aspire. Act. Achieve.", font_size=18, color=CYAN, alignment=PP_ALIGN.CENTER)

    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(3.5), Inches(3.9), Inches(3), Inches(0.04),
              fill_color=RED)

    add_text_box(slide, Inches(1), Inches(4.2), Inches(8), Inches(1.5),
                 "Jack Mlusu, Founder & CEO\n"
                 "jmlusu@gmail.com\n"
                 "+265 (0) 980 016 004\n"
                 "lightspeedholdings.com",
                 font_size=14, color=WHITE, alignment=PP_ALIGN.CENTER)

    add_text_box(slide, Inches(1), Inches(6.3), Inches(8), Inches(0.5),
                 "Open by design. Governed by construction.",
                 font_size=12, color=CYAN, alignment=PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    create_title_slide(prs)
    create_purpose_slide(prs)
    create_policy_slide(prs)
    create_who_we_are_slide(prs)
    create_delivery_slide(prs)
    create_offers_slide(prs)
    create_potential_slide(prs)
    create_governance_slide(prs)
    create_data_protection_slide(prs)
    create_transparency_slide(prs)
    create_audit_slide(prs)
    create_alignment_slide(prs)
    create_regulator_slide(prs)
    create_thank_you_slide(prs)

    output_path = os.path.join(os.path.dirname(__file__), "lightspeed-regulator-deck.pptx")
    prs.save(output_path)
    print(f"Malawi regulator deck saved to: {os.path.abspath(output_path)}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()