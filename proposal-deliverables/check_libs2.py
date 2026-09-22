import docx

# Extract text from VillageReach USSD docx
print("=== VillageReach-USSD-Consultant-Proposal.docx ===")
doc = docx.Document("VillageReach-USSD-Consultant-Proposal.docx")
full_text = "\n".join([para.text for para in doc.paragraphs])
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Text length: {len(full_text)}")
# Search for key terms
for s in [
    "1,875",
    "1875",
    "2024",
    "Last updated",
    "$45,000",
    "45000",
    "54747",
    "Malawi",
    "Data Protection",
    "Jack Mlusu",
    "no callback",
    "no-callback",
]:
    found = s in full_text
    print(f"  '{s}': {'FOUND' if found else 'MISSING'}")

# Also search in tables
print("\n--- Searching table cells ---")
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if s := "54747" in cell.text:
                print(f"Found 54747 in table cell: {cell.text[:50]}")

print()

# Extract text from Jack-Mlusu docx
print("=== Jack-Mlusu-Solo-Consulting-Proposal.docx ===")
doc2 = docx.Document(
    r"C:\Users\jmlus\light-speed-holdings\Jack-Mlusu-Solo-Consulting-Proposal.docx"
)
full_text2 = "\n".join([para.text for para in doc2.paragraphs])
print(f"Total paragraphs: {len(doc2.paragraphs)}")
print(f"Text length: {len(full_text2)}")
for s in [
    "1,875",
    "1875",
    "2024",
    "Last updated",
    "$45,000",
    "45000",
    "54747",
    "Malawi",
    "Data Protection",
    "Jack Mlusu",
    "no callback",
    "no-callback",
    "Malawi Data Protection Act",
]:
    found = s in full_text2
    print(f"  '{s}': {'FOUND' if found else 'MISSING'}")

print()

# Extract text from Jack-Mlusu pptx
print("=== Jack-Mlusu-Solo-Consultant-Proposal-Slides.pptx ===")
from pptx import Presentation  # noqa: E402

prs = Presentation(
    r"C:\Users\jmlus\light-speed-holdings\Jack-Mlusu-Solo-Consultant-Proposal-Slides.pptx"
)
full_text3 = ""
for slide in prs.slides:
    for shape in slide.shapes:
        if shape.has_text_frame:
            full_text3 += shape.text + "\n"
print(f"Total slides: {len(prs.slides)}")
print(f"Text length: {len(full_text3)}")
for s in [
    "1,875",
    "1875",
    "2024",
    "Last updated",
    "$45,000",
    "45000",
    "54747",
    "Malawi",
    "Data Protection",
    "Jack Mlusu",
    "no callback",
    "no-callback",
    "Malawi Data Protection Act",
]:
    found = s in full_text3
    print(f"  '{s}': {'FOUND' if found else 'MISSING'}")
